from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from pydantic import ValidationError

from app.schemas import Node, Order, Scenario, Vehicle

MAX_IMPORT_BYTES = 5 * 1024 * 1024
REQUIRED_SHEETS = {"场景", "节点", "订单", "车辆"}
SHEET_HEADERS = {
    "场景": ["scenario_code", "name", "county_name"],
    "节点": ["id", "name", "longitude", "latitude", "kind"],
    "订单": [
        "id",
        "node_id",
        "weight_kg",
        "window_start",
        "window_end",
        "service_minutes",
        "cold_chain",
        "priority",
    ],
    "车辆": [
        "id",
        "plate",
        "capacity_kg",
        "cold_chain",
        "available",
        "cost_per_km",
        "shift_start",
        "shift_end",
    ],
}


class ScenarioImportError(ValueError):
    def __init__(self, errors: list[str]) -> None:
        super().__init__("；".join(errors))
        self.errors = errors


def _rows(worksheet: Any, expected_headers: list[str]) -> list[dict[str, Any]]:
    values = list(worksheet.iter_rows(values_only=True))
    if not values:
        raise ScenarioImportError([f"工作表“{worksheet.title}”为空"])
    headers = [str(value).strip() if value is not None else "" for value in values[0]]
    missing = [header for header in expected_headers if header not in headers]
    if missing:
        raise ScenarioImportError(
            [f"工作表“{worksheet.title}”缺少字段：{', '.join(missing)}"]
        )
    indexes = {header: headers.index(header) for header in expected_headers}
    result: list[dict[str, Any]] = []
    for row_number, row in enumerate(values[1:], start=2):
        if all(value is None or str(value).strip() == "" for value in row):
            continue
        item = {
            header: row[indexes[header]] if indexes[header] < len(row) else None
            for header in expected_headers
        }
        item["_row"] = row_number
        result.append(item)
    return result


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _boolean(value: Any, *, sheet: str, row: int, field: str) -> bool:
    if isinstance(value, bool):
        return value
    normalized = _text(value).lower()
    if normalized in {"是", "true", "1", "yes", "y"}:
        return True
    if normalized in {"否", "false", "0", "no", "n", ""}:
        return False
    raise ScenarioImportError([f"{sheet}第 {row} 行字段 {field} 必须填写“是”或“否”"])


def parse_scenario_workbook(content: bytes) -> tuple[str, str, Scenario, list[str]]:
    if len(content) > MAX_IMPORT_BYTES:
        raise ScenarioImportError(["文件不能超过 5 MB"])
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ScenarioImportError(["无法读取 Excel 文件，请使用系统提供的 .xlsx 模板"]) from exc

    missing_sheets = sorted(REQUIRED_SHEETS.difference(workbook.sheetnames))
    if missing_sheets:
        raise ScenarioImportError([f"缺少工作表：{', '.join(missing_sheets)}"])

    scenario_rows = _rows(workbook["场景"], SHEET_HEADERS["场景"])
    if len(scenario_rows) != 1:
        raise ScenarioImportError(["“场景”工作表必须且只能包含一条数据"])
    scenario_row = scenario_rows[0]
    scenario_code = _text(scenario_row["scenario_code"])
    name = _text(scenario_row["name"])
    county_name = _text(scenario_row["county_name"]) or "祁东县"
    if not scenario_code or not name:
        raise ScenarioImportError(["场景代码和场景名称不能为空"])

    errors: list[str] = []
    nodes: list[Node] = []
    for item in _rows(workbook["节点"], SHEET_HEADERS["节点"]):
        try:
            nodes.append(
                Node(
                    id=_text(item["id"]),
                    name=_text(item["name"]),
                    longitude=item["longitude"],
                    latitude=item["latitude"],
                    kind=_text(item["kind"]),
                )
            )
        except (ValidationError, TypeError, ValueError) as exc:
            errors.append(f"节点第 {item['_row']} 行：{exc}")

    orders: list[Order] = []
    for item in _rows(workbook["订单"], SHEET_HEADERS["订单"]):
        try:
            orders.append(
                Order(
                    id=_text(item["id"]),
                    node_id=_text(item["node_id"]),
                    weight_kg=item["weight_kg"],
                    window_start=item["window_start"],
                    window_end=item["window_end"],
                    service_minutes=item["service_minutes"] or 8,
                    cold_chain=_boolean(
                        item["cold_chain"],
                        sheet="订单",
                        row=item["_row"],
                        field="cold_chain",
                    ),
                    priority=_text(item["priority"]) or "normal",
                )
            )
        except (ScenarioImportError, ValidationError, TypeError, ValueError) as exc:
            errors.append(f"订单第 {item['_row']} 行：{exc}")

    vehicles: list[Vehicle] = []
    for item in _rows(workbook["车辆"], SHEET_HEADERS["车辆"]):
        try:
            vehicles.append(
                Vehicle(
                    id=_text(item["id"]),
                    plate=_text(item["plate"]),
                    capacity_kg=item["capacity_kg"],
                    cold_chain=_boolean(
                        item["cold_chain"],
                        sheet="车辆",
                        row=item["_row"],
                        field="cold_chain",
                    ),
                    available=_boolean(
                        item["available"],
                        sheet="车辆",
                        row=item["_row"],
                        field="available",
                    ),
                    cost_per_km=item["cost_per_km"],
                    shift_start=item["shift_start"] or 0,
                    shift_end=item["shift_end"] or 600,
                )
            )
        except (ScenarioImportError, ValidationError, TypeError, ValueError) as exc:
            errors.append(f"车辆第 {item['_row']} 行：{exc}")

    if errors:
        raise ScenarioImportError(errors[:30])
    validation_errors, warnings = validate_scenario(nodes, orders, vehicles)
    if validation_errors:
        raise ScenarioImportError(validation_errors)

    scenario = Scenario(
        id=scenario_code,
        name=name,
        nodes=nodes,
        orders=orders,
        vehicles=vehicles,
    )
    return scenario_code, county_name, scenario, warnings


def validate_scenario(
    nodes: list[Node],
    orders: list[Order],
    vehicles: list[Vehicle],
) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    if not nodes:
        errors.append("至少需要一个物流节点")
    if not orders:
        errors.append("至少需要一个订单")
    if not vehicles:
        errors.append("至少需要一台车辆")

    for label, values in (
        ("节点", [item.id for item in nodes]),
        ("订单", [item.id for item in orders]),
        ("车辆", [item.id for item in vehicles]),
        ("车牌", [item.plate for item in vehicles]),
    ):
        duplicates = sorted({value for value in values if values.count(value) > 1})
        if duplicates:
            errors.append(f"{label}存在重复值：{', '.join(duplicates)}")

    depots = [node for node in nodes if node.kind == "depot"]
    if len(depots) != 1:
        errors.append("必须且只能配置一个 depot 类型的共配中心")
    node_ids = {node.id for node in nodes}
    unknown_nodes = sorted({order.node_id for order in orders if order.node_id not in node_ids})
    if unknown_nodes:
        errors.append(f"订单引用了不存在的节点：{', '.join(unknown_nodes)}")
    if any(order.cold_chain for order in orders) and not any(
        vehicle.available and vehicle.cold_chain for vehicle in vehicles
    ):
        errors.append("存在冷链订单，但没有可用冷链车辆")
    if sum(order.weight_kg for order in orders) > sum(
        vehicle.capacity_kg for vehicle in vehicles if vehicle.available
    ) * 3:
        warnings.append("订单总重量明显高于单轮可用运力，可能出现未分配订单")
    return errors, warnings
