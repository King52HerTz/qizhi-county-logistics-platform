from __future__ import annotations

import csv
import hashlib
import json
from datetime import datetime, time, timezone
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import ValidationError
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import (
    DataImportBatchRecord,
    DataQualityRuleRecord,
    NodeRecord,
    OrderRecord,
    VehicleRecord,
)
from app.schemas import MasterOrderCreate, Node, Order, Scenario, Vehicle
from app.spatial import gcj02_to_wgs84, inside_qidong_bounds


MAX_ORDER_IMPORT_BYTES = 5 * 1024 * 1024
ORDER_IMPORT_HEADERS = [
    "订单编号",
    "货物名称",
    "客户名称",
    "联系电话",
    "服务网点编号",
    "配送地址",
    "重量(kg)",
    "时间窗开始",
    "时间窗结束",
    "服务时长(分钟)",
    "冷链要求",
    "优先级",
]
HEADER_FIELDS = {
    "订单编号": "id",
    "货物名称": "cargo_name",
    "客户名称": "customer_name",
    "联系电话": "customer_phone",
    "服务网点编号": "node_id",
    "配送地址": "delivery_address",
    "重量(kg)": "weight_kg",
    "时间窗开始": "window_start",
    "时间窗结束": "window_end",
    "服务时长(分钟)": "service_minutes",
    "冷链要求": "cold_chain",
    "优先级": "priority",
}


class OrderImportError(ValueError):
    pass


def build_order_import_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "订单导入"
    sheet.append(ORDER_IMPORT_HEADERS)
    sheet.append(
        [
            "QD20260726001",
            "黄花菜",
            "祁东示范合作社",
            "13800000000",
            "V-YINXING",
            "祁东县银星村",
            680,
            "08:00",
            "12:00",
            8,
            "否",
            "普通",
        ]
    )
    header_fill = PatternFill("solid", fgColor="0F6A45")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    widths = [18, 16, 22, 17, 19, 28, 13, 15, 15, 18, 13, 13]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:L2"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _clean_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _boolean(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    normalized = _clean_text(value).lower()
    if normalized in {"是", "true", "1", "yes", "y"}:
        return True
    if normalized in {"否", "false", "0", "no", "n", ""}:
        return False
    raise ValueError("必须填写“是”或“否”")


def _priority(value: Any) -> str:
    normalized = _clean_text(value).lower()
    if normalized in {"普通", "normal", ""}:
        return "normal"
    if normalized in {"紧急", "urgent"}:
        return "urgent"
    raise ValueError("只能填写“普通”或“紧急”")


def _minute(value: Any) -> int:
    if isinstance(value, time):
        return value.hour * 60 + value.minute
    if isinstance(value, datetime):
        return value.hour * 60 + value.minute
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, float) and 0 <= value < 1:
            return round(value * 24 * 60)
        return int(value)
    normalized = _clean_text(value)
    if ":" in normalized:
        hour, minute = normalized.split(":", 1)
        result = int(hour) * 60 + int(minute)
    else:
        result = int(normalized)
    if not 0 <= result <= 1440:
        raise ValueError("必须为 00:00 至 24:00")
    return result


def _input_rows(filename: str, content: bytes) -> list[tuple[int, dict[str, Any]]]:
    if len(content) > MAX_ORDER_IMPORT_BYTES:
        raise OrderImportError("文件不能超过 5 MB")
    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx":
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            sheet = workbook["订单导入"] if "订单导入" in workbook.sheetnames else workbook.active
            values = list(sheet.iter_rows(values_only=True))
        except Exception as exc:
            raise OrderImportError("无法读取 Excel 文件，请使用系统订单模板") from exc
        if not values:
            raise OrderImportError("导入文件为空")
        headers = [_clean_text(value) for value in values[0]]
        missing = [header for header in ORDER_IMPORT_HEADERS if header not in headers]
        if missing:
            raise OrderImportError(f"缺少字段：{', '.join(missing)}")
        indexes = {header: headers.index(header) for header in ORDER_IMPORT_HEADERS}
        result: list[tuple[int, dict[str, Any]]] = []
        for row_number, row in enumerate(values[1:], start=2):
            if all(value is None or _clean_text(value) == "" for value in row):
                continue
            result.append(
                (
                    row_number,
                    {
                        HEADER_FIELDS[header]: (
                            row[indexes[header]] if indexes[header] < len(row) else None
                        )
                        for header in ORDER_IMPORT_HEADERS
                    },
                )
            )
        return result
    if suffix == ".csv":
        try:
            decoded = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise OrderImportError("CSV 文件必须使用 UTF-8 编码") from exc
        reader = csv.DictReader(StringIO(decoded))
        missing = [header for header in ORDER_IMPORT_HEADERS if header not in (reader.fieldnames or [])]
        if missing:
            raise OrderImportError(f"缺少字段：{', '.join(missing)}")
        return [
            (
                row_number,
                {HEADER_FIELDS[header]: row.get(header) for header in ORDER_IMPORT_HEADERS},
            )
            for row_number, row in enumerate(reader, start=2)
            if any(_clean_text(value) for value in row.values())
        ]
    raise OrderImportError("只支持 .xlsx 或 UTF-8 CSV 订单文件")


def _validation_message(error: ValidationError) -> tuple[str, str]:
    first = error.errors()[0]
    field = str(first.get("loc", ["数据"])[-1])
    names = {
        "id": "订单编号",
        "node_id": "服务网点编号",
        "cargo_name": "货物名称",
        "customer_name": "客户名称",
        "weight_kg": "重量(kg)",
        "window_start": "时间窗开始",
        "window_end": "时间窗结束",
        "service_minutes": "服务时长(分钟)",
        "priority": "优先级",
    }
    message = str(first.get("msg", "字段格式不正确"))
    if first.get("type") == "value_error":
        message = message.removeprefix("Value error, ")
    return names.get(field, field), message


def preview_order_import(
    session: Session,
    *,
    filename: str,
    content: bytes,
    created_by: int,
) -> DataImportBatchRecord:
    rows = _input_rows(filename, content)
    active_nodes = set(
        session.scalars(select(NodeRecord.id).where(NodeRecord.is_active.is_(True))).all()
    )
    existing_ids = set(session.scalars(select(OrderRecord.id)).all())
    seen_ids: set[str] = set()
    payload: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []

    for row_number, raw in rows:
        normalized: dict[str, Any]
        try:
            normalized = {
                "id": _clean_text(raw["id"]),
                "cargo_name": _clean_text(raw["cargo_name"]),
                "customer_name": _clean_text(raw["customer_name"]),
                "customer_phone": _clean_text(raw["customer_phone"]) or None,
                "node_id": _clean_text(raw["node_id"]),
                "delivery_address": _clean_text(raw["delivery_address"]) or None,
                "weight_kg": int(raw["weight_kg"]),
                "window_start": _minute(raw["window_start"]),
                "window_end": _minute(raw["window_end"]),
                "service_minutes": int(raw["service_minutes"] or 8),
                "cold_chain": _boolean(raw["cold_chain"]),
                "priority": _priority(raw["priority"]),
                "status": "pending",
            }
            validated = MasterOrderCreate.model_validate(normalized)
        except ValidationError as exc:
            field, message = _validation_message(exc)
            errors.append({"row": row_number, "field": field, "message": message})
            continue
        except (TypeError, ValueError) as exc:
            errors.append({"row": row_number, "field": "数据格式", "message": str(exc)})
            continue

        if validated.id in existing_ids:
            errors.append(
                {"row": row_number, "field": "订单编号", "message": "订单编号已存在"}
            )
            continue
        if validated.id in seen_ids:
            errors.append(
                {"row": row_number, "field": "订单编号", "message": "文件内订单编号重复"}
            )
            continue
        seen_ids.add(validated.id)
        if validated.node_id not in active_nodes:
            errors.append(
                {
                    "row": row_number,
                    "field": "服务网点编号",
                    "message": "网点不存在或未启用",
                }
            )
            continue
        payload.append(validated.model_dump(mode="json"))

    if not rows:
        errors.append({"row": 0, "field": "文件", "message": "没有可导入的数据行"})
    error_row_count = len({item["row"] for item in errors})
    record = DataImportBatchRecord(
        batch_code=f"IMP{datetime.now(timezone.utc):%Y%m%d}{uuid4().hex[:6].upper()}",
        import_type="orders",
        filename=Path(filename).name,
        status="rejected" if errors else "validated",
        total_rows=len(rows),
        valid_rows=len(payload),
        error_rows=error_row_count,
        warning_rows=0,
        payload=payload,
        errors=errors[:200],
        created_by=created_by,
    )
    session.add(record)
    session.commit()
    session.refresh(record)
    return record


def commit_order_import(
    session: Session,
    record: DataImportBatchRecord,
    *,
    committed_by: int,
) -> DataImportBatchRecord:
    if record.status == "committed":
        raise OrderImportError("该导入批次已经入库")
    if record.status != "validated" or record.error_rows:
        raise OrderImportError("批次存在阻断错误，修正文件后请重新预检")
    order_ids = [item["id"] for item in record.payload]
    conflicts = set(
        session.scalars(select(OrderRecord.id).where(OrderRecord.id.in_(order_ids))).all()
    )
    if conflicts:
        raise OrderImportError(f"订单编号已被其他操作占用：{', '.join(sorted(conflicts))}")
    active_nodes = set(
        session.scalars(select(NodeRecord.id).where(NodeRecord.is_active.is_(True))).all()
    )
    invalid_nodes = sorted(
        {item["node_id"] for item in record.payload if item["node_id"] not in active_nodes}
    )
    if invalid_nodes:
        raise OrderImportError(f"网点已停用或删除：{', '.join(invalid_nodes)}")

    session.add_all([OrderRecord(**item) for item in record.payload])
    record.status = "committed"
    record.committed_by = committed_by
    record.committed_at = datetime.now(timezone.utc)
    session.commit()
    session.refresh(record)
    return record


def quality_report(session: Session) -> list[dict[str, Any]]:
    rules = list(
        session.scalars(select(DataQualityRuleRecord).order_by(DataQualityRuleRecord.id))
    )
    pending_orders = session.scalar(
        select(func.count(OrderRecord.id)).where(OrderRecord.status.in_(["pending", "planned"]))
    ) or 0
    inactive_node_orders = session.scalar(
        select(func.count(OrderRecord.id))
        .join(NodeRecord, OrderRecord.node_id == NodeRecord.id)
        .where(
            OrderRecord.status.in_(["pending", "planned"]),
            NodeRecord.is_active.is_(False),
        )
    ) or 0
    invalid_windows = session.scalar(
        select(func.count(OrderRecord.id)).where(
            OrderRecord.status.in_(["pending", "planned"]),
            OrderRecord.window_end <= OrderRecord.window_start,
        )
    ) or 0
    order_weight = session.scalar(
        select(func.coalesce(func.sum(OrderRecord.weight_kg), 0)).where(
            OrderRecord.status.in_(["pending", "planned"])
        )
    ) or 0
    vehicle_capacity = session.scalar(
        select(func.coalesce(func.sum(VehicleRecord.capacity_kg), 0)).where(
            VehicleRecord.status == "available"
        )
    ) or 0
    cold_orders = session.scalar(
        select(func.count(OrderRecord.id)).where(
            OrderRecord.status.in_(["pending", "planned"]),
            OrderRecord.cold_chain.is_(True),
        )
    ) or 0
    cold_vehicles = session.scalar(
        select(func.count(VehicleRecord.id)).where(
            VehicleRecord.status == "available",
            VehicleRecord.cold_chain.is_(True),
        )
    ) or 0
    depot_count = session.scalar(
        select(func.count(NodeRecord.id)).where(
            NodeRecord.kind == "depot",
            NodeRecord.is_active.is_(True),
        )
    ) or 0

    issue_counts = {
        "order_active_node": inactive_node_orders,
        "order_time_window": invalid_windows,
        "vehicle_capacity": max(
            0,
            int(order_weight) - int(vehicle_capacity) * 3,
        ),
        "cold_chain_vehicle": cold_orders if cold_orders and not cold_vehicles else 0,
        "single_depot": 0 if depot_count == 1 else abs(depot_count - 1) or 1,
    }
    details = {
        "order_active_node": f"{inactive_node_orders} 个订单引用停用网点",
        "order_time_window": f"{invalid_windows} 个订单时间窗异常",
        "vehicle_capacity": (
            f"待调度 {int(order_weight)} kg / 三轮可用运力 {int(vehicle_capacity) * 3} kg"
        ),
        "cold_chain_vehicle": f"{cold_orders} 个冷链订单 / {cold_vehicles} 台可用冷链车",
        "single_depot": f"当前启用共配中心 {depot_count} 个",
    }
    result: list[dict[str, Any]] = []
    for rule in rules:
        count = int(issue_counts.get(rule.rule_code, 0))
        if not rule.enabled:
            rule_status = "disabled"
        elif count == 0:
            rule_status = "passed"
        elif rule.severity == "warning":
            rule_status = "warning"
        else:
            rule_status = "blocked"
        result.append(
            {
                "id": rule.id,
                "rule_code": rule.rule_code,
                "name": rule.name,
                "description": rule.description,
                "entity_type": rule.entity_type,
                "severity": rule.severity,
                "enabled": rule.enabled,
                "config": rule.config,
                "issue_count": count,
                "status": rule_status,
                "detail": details.get(rule.rule_code, f"待发布订单 {pending_orders} 个"),
                "updated_at": rule.updated_at,
            }
        )
    return result


def build_scenario_from_master(session: Session, *, name: str) -> tuple[Scenario, list[str]]:
    report = quality_report(session)
    blockers = [item for item in report if item["status"] == "blocked"]
    if blockers:
        raise OrderImportError(
            "发布检查未通过：" + "；".join(item["name"] for item in blockers)
        )
    nodes = list(
        session.scalars(
            select(NodeRecord)
            .where(NodeRecord.is_active.is_(True))
            .order_by(NodeRecord.id)
        )
    )
    orders = list(
        session.scalars(
            select(OrderRecord)
            .where(OrderRecord.status != "cancelled")
            .order_by(OrderRecord.id)
        )
    )
    vehicles = list(
        session.scalars(
            select(VehicleRecord)
            .where(VehicleRecord.status != "offline")
            .order_by(VehicleRecord.id)
        )
    )
    if not orders:
        raise OrderImportError("没有可发布的订单")
    if not any(vehicle.status == "available" for vehicle in vehicles):
        raise OrderImportError("没有可用车辆")
    node_ids = {record.id for record in nodes}
    if any(record.node_id not in node_ids for record in orders):
        raise OrderImportError("存在引用停用网点的订单")
    depots = [record for record in nodes if record.kind == "depot"]
    if len(depots) != 1:
        raise OrderImportError(f"必须且只能启用一个配送中心，当前为 {len(depots)} 个")
    outside_nodes = [
        record.id
        for record in nodes
        if not inside_qidong_bounds(float(record.longitude), float(record.latitude))
    ]
    if outside_nodes:
        raise OrderImportError(
            "存在超出祁东县业务范围的网点：" + "、".join(outside_nodes[:8])
        )
    nodes.sort(key=lambda record: (record.kind != "depot", record.id))

    snapshot_items = [
        {
            "id": record.id,
            "longitude": round(float(record.longitude), 6),
            "latitude": round(float(record.latitude), 6),
            "status": record.coordinate_status,
            "calibrated_at": (
                record.calibrated_at.isoformat() if record.calibrated_at else None
            ),
        }
        for record in nodes
    ]
    snapshot_id = "spatial-" + hashlib.sha256(
        json.dumps(snapshot_items, ensure_ascii=False, sort_keys=True).encode("utf-8")
    ).hexdigest()[:16]
    snapshot_at = datetime.now(timezone.utc)
    verified_nodes = sum(
        record.coordinate_status == "verified" for record in nodes
    )

    scenario = Scenario(
        id="master-data",
        name=name,
        spatial_snapshot_id=snapshot_id,
        spatial_snapshot_at=snapshot_at,
        spatial_verified_nodes=verified_nodes,
        nodes=[
            Node(
                id=record.id,
                name=record.name,
                longitude=float(record.longitude),
                latitude=float(record.latitude),
                kind=record.kind,
                coordinate_system="gcj02",
                coordinate_source=record.coordinate_source,
                coordinate_status=record.coordinate_status,
                wgs84_longitude=(
                    float(record.wgs84_longitude)
                    if record.wgs84_longitude is not None
                    else gcj02_to_wgs84(
                        float(record.longitude), float(record.latitude)
                    )[0]
                ),
                wgs84_latitude=(
                    float(record.wgs84_latitude)
                    if record.wgs84_latitude is not None
                    else gcj02_to_wgs84(
                        float(record.longitude), float(record.latitude)
                    )[1]
                ),
                coordinate_accuracy_m=record.coordinate_accuracy_m,
                calibrated_at=record.calibrated_at,
            )
            for record in nodes
        ],
        orders=[
            Order(
                id=record.id,
                node_id=record.node_id,
                weight_kg=record.weight_kg,
                window_start=record.window_start,
                window_end=record.window_end,
                service_minutes=record.service_minutes,
                cold_chain=record.cold_chain,
                priority=record.priority,
                completed=record.status == "completed",
            )
            for record in orders
        ],
        vehicles=[
            Vehicle(
                id=record.id,
                plate=record.plate,
                capacity_kg=record.capacity_kg,
                cold_chain=record.cold_chain,
                available=record.status == "available",
                cost_per_km=float(record.cost_per_km),
                shift_start=record.shift_start,
                shift_end=record.shift_end,
            )
            for record in vehicles
        ],
    )
    warnings = [
        item["detail"] for item in report if item["status"] == "warning"
    ]
    if verified_nodes < len(nodes):
        warnings.append(
            f"空间坐标快照包含 {len(nodes) - verified_nodes} 个未完成核验的网点；"
            "建议完成校准后重新发布"
        )
    return scenario, warnings
