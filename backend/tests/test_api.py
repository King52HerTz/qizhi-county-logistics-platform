from io import BytesIO
import json
from pathlib import Path
from uuid import uuid4

import pytest
import httpx
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.ai_service import _plain_text
from app.db import SessionLocal
from app.main import app
from app.optimization_orchestrator import orchestrate_optimization
from app.models import (
    AIProviderConfigRecord,
    DataImportBatchRecord,
    EventRecord,
    MapProviderConfigRecord,
    NodeRecord,
    OrderRecord,
    ScenarioVersionRecord,
)
from app.config import get_settings
from app.routing_matrix import build_routing_matrix, clear_matrix_cache
from app.data.scenario import get_scenario
from app.solver import solve_scenario
from app.spatial import gcj02_to_wgs84, haversine_m, inside_qidong_bounds
from app.schemas import CurrentWeather, OptimizeRequest


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as test_client:
        login = test_client.post(
            "/api/v1/auth/login",
            json={"username": "admin", "password": "QizhiDemo2026!"},
        )
        assert login.status_code == 200
        test_client.headers.update(
            {"Authorization": f"Bearer {login.json()['access_token']}"}
        )
        yield test_client


def test_health_reports_database(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    response = client.get("/api/v1/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"
    assert response.json()["database"] is True
    live = client.get("/api/v1/health/live")
    assert live.status_code == 200
    assert live.json()["status"] == "ok"

    class RedisProbe:
        def ping(self) -> bool:
            return True

        def close(self) -> None:
            return None

    monkeypatch.setattr("app.main.redis.Redis.from_url", lambda *args, **kwargs: RedisProbe())
    ready = client.get("/api/v1/health/ready")
    assert ready.status_code == 200
    assert ready.json()["checks"] == {
        "database": True,
        "redis": True,
        "migrations": True,
    }
    openapi_paths = client.get("/api/v1/openapi.json").json()["paths"]
    assert "/api/v1/operations/readiness" in openapi_paths
    assert "/api/v1/competition/readiness" not in openapi_paths


def test_spatial_coordinate_math_uses_explicit_crs() -> None:
    wgs_longitude, wgs_latitude = gcj02_to_wgs84(112.08, 26.78)
    assert abs(wgs_longitude - 112.08) > 0.001
    assert abs(wgs_latitude - 26.78) > 0.001
    assert inside_qidong_bounds(112.08, 26.78) is True
    assert inside_qidong_bounds(120.0, 30.0) is False
    assert 900 < haversine_m(112.08, 26.78, 112.09, 26.78) < 1100


def test_forecast_and_routing_status_contract(client: TestClient) -> None:
    forecast = client.get("/api/v1/analytics/forecast?days=7")
    assert forecast.status_code == 200
    payload = forecast.json()
    assert payload["horizon_days"] == 7
    assert len([point for point in payload["points"] if point["actual_tons"] is not None]) == payload["training_samples"]
    assert len([point for point in payload["points"] if point["forecast_tons"] is not None]) == 7
    assert payload["township_ranking"]
    assert payload["recommended_vehicles"] >= 1
    assert payload["risk_level"] in {"low", "medium", "high"}
    assert payload["fallback_used"] is True
    assert payload["model_family"] == "deterministic-seasonal-baseline"
    assert payload["fallback_reason"]
    assert client.get("/api/v1/analytics/forecast?days=2").status_code == 422

    status_response = client.get("/api/v1/routing/status")
    assert status_response.status_code == 200
    assert status_response.json()["fallback_enabled"] is True


def test_current_weather_contract(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    async def fake_weather() -> CurrentWeather:
        return CurrentWeather(
            location="祁东县",
            temperature_c=31.6,
            apparent_temperature_c=34.2,
            humidity_percent=68,
            weather_code=2,
            weather_label="少云",
            wind_speed_kmh=7.4,
            observed_at="2026-07-29T20:15",
            fetched_at="2026-07-29T12:16:00Z",
            source="Open-Meteo",
        )

    monkeypatch.setattr("app.main.get_current_weather", fake_weather)
    response = client.get("/api/v1/weather/current")
    assert response.status_code == 200
    assert response.json()["location"] == "祁东县"
    assert response.json()["temperature_c"] == 31.6
    assert response.json()["weather_label"] == "少云"


def test_osrm_matrix_and_memory_cache(monkeypatch: pytest.MonkeyPatch) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "routing_provider", "osrm")
    clear_matrix_cache()
    calls = 0

    def fake_get(*_: object, **__: object) -> httpx.Response:
        nonlocal calls
        calls += 1
        request = httpx.Request("GET", "https://router.test/table")
        return httpx.Response(
            200,
            request=request,
            json={
                "code": "Ok",
                "data_version": "2026-07-20T00:00:00Z",
                "distances": [[0, 1240], [1320, 0]],
                "durations": [[0, 180], [240, 0]],
            },
        )

    monkeypatch.setattr("app.routing_matrix.httpx.get", fake_get)
    coords = [(112.1, 26.8), (112.2, 26.9)]
    first = build_routing_matrix(coords)
    second = build_routing_matrix(coords)
    assert first.info.source == "osrm"
    assert first.info.cache_hit is False
    assert first.duration_min == [[0, 3], [4, 0]]
    assert second.info.cache_hit is True
    assert calls == 1
    clear_matrix_cache()


def test_osrm_failure_falls_back_without_interrupting_solver(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "routing_provider", "osrm")
    clear_matrix_cache()

    def failed_get(*_: object, **__: object) -> httpx.Response:
        raise httpx.ConnectError("OSRM unavailable")

    monkeypatch.setattr("app.routing_matrix.httpx.get", failed_get)
    matrix = build_routing_matrix([(112.1, 26.8), (112.2, 26.9)])
    assert matrix.info.source == "haversine"
    assert matrix.info.fallback_used is True
    assert "OSRM unavailable" in (matrix.info.fallback_reason or "")
    assert matrix.distance_m[0][1] > 0
    clear_matrix_cache()


def test_standard_profile_uses_fixed_snapshot_and_is_reproducible(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    snapshot = json.loads(
        (
            Path(__file__).parents[1]
            / "app"
            / "data"
            / "standard_demo_matrix.json"
        ).read_text(encoding="utf-8")
    )
    scenario = get_scenario()
    service_node_ids = [
        next(node.id for node in scenario.nodes if node.kind == "depot"),
        *[order.node_id for order in scenario.orders],
    ]
    coordinates = [
        tuple(map(float, value.split(",")))
        for value in snapshot["base"]["coordinates"]
    ]
    coordinate_by_node = dict(zip(service_node_ids, coordinates, strict=True))
    scenario.nodes = [
        node.model_copy(
            update={
                "longitude": coordinate_by_node[node.id][0],
                "latitude": coordinate_by_node[node.id][1],
            }
        )
        if node.id in coordinate_by_node
        else node
        for node in scenario.nodes
    ]
    calls = 0

    def unexpected_get(*_: object, **__: object) -> httpx.Response:
        nonlocal calls
        calls += 1
        raise AssertionError("标准运营基线不应请求实时 OSRM")

    monkeypatch.setattr("app.routing_matrix.httpx.get", unexpected_get)
    clear_matrix_cache()
    results = [
        orchestrate_optimization(
            scenario,
            OptimizeRequest(
                profile="standard-v1",
                strategy="balanced",
                ai_assisted=True,
                candidate_count=1,
            ),
        )
        for _ in range(3)
    ]
    canonical = [
        {
            "routes": [route.model_dump() for route in result.routes],
            "unassigned": result.unassigned_order_ids,
            "metrics": result.metrics.model_dump(),
            "baseline": result.baseline_metrics.model_dump(),
            "data_version": result.routing.data_version,
            "candidates": [
                candidate.model_dump() for candidate in result.candidate_summaries
            ],
        }
        for result in results
    ]
    assert canonical[0] == canonical[1] == canonical[2]
    assert results[0].routing.source == "osrm"
    assert results[0].routing.data_version == "standard-v1-20260729"
    assert results[0].ai_orchestration.used is False
    assert calls == 0
    clear_matrix_cache()


def test_ai_explanation_is_normalized_to_plain_text() -> None:
    assert _plain_text("## **里程**\n`302.2 km`") == "里程\n302.2 km"


def test_admin_login_and_current_user(client: TestClient) -> None:
    response = client.post(
        "/api/v1/auth/login",
        json={"username": "admin", "password": "QizhiDemo2026!"},
    )
    assert response.status_code == 200
    token = response.json()["access_token"]
    profile = client.get("/api/v1/auth/me", headers={"Authorization": f"Bearer {token}"})
    assert profile.status_code == 200
    assert profile.json()["username"] == "admin"
    assert profile.json()["role"] == "admin"
    assert profile.json()["display_name"]
    assert profile.json()["last_login_at"]


def test_public_registration_creates_least_privilege_account(client: TestClient) -> None:
    username = f"user_{uuid4().hex[:10]}"
    response = client.post(
        "/api/v1/auth/register",
        json={
            "username": username,
            "display_name": "新注册用户",
            "password": "StrongPass2026!",
            "agreement_accepted": True,
        },
    )
    assert response.status_code == 201
    assert response.json()["username"] == username
    assert response.json()["display_name"] == "新注册用户"
    assert response.json()["role"] == "driver"

    profile = client.get(
        "/api/v1/auth/me",
        headers={"Authorization": f"Bearer {response.json()['access_token']}"},
    )
    assert profile.status_code == 200
    assert profile.json()["role"] == "driver"


def test_public_registration_rejects_duplicate_weak_and_privileged_payloads(
    client: TestClient,
) -> None:
    username = f"user_{uuid4().hex[:10]}"
    valid_payload = {
        "username": username,
        "display_name": "安全测试用户",
        "password": "StrongPass2026!",
        "agreement_accepted": True,
    }
    assert client.post("/api/v1/auth/register", json=valid_payload).status_code == 201
    duplicate = client.post("/api/v1/auth/register", json=valid_payload)
    assert duplicate.status_code == 409

    weak = client.post(
        "/api/v1/auth/register",
        json={**valid_payload, "username": f"user_{uuid4().hex[:10]}", "password": "password"},
    )
    assert weak.status_code == 422

    privileged = client.post(
        "/api/v1/auth/register",
        json={**valid_payload, "username": f"user_{uuid4().hex[:10]}", "role": "admin"},
    )
    assert privileged.status_code == 422


def test_optimization_respects_cold_chain_and_persists(client: TestClient) -> None:
    response = client.post(
        "/api/v1/optimize",
        json={
            "scenario_id": "qidong-demo-001",
            "solver": "ortools",
            "time_limit_seconds": 1,
            "options": {"cold_chain_first": True, "backhaul_enabled": True},
        },
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] in {"optimal", "feasible"}
    assert payload["routes"]
    assert 0 <= payload["decision_score"]["total"] <= 100
    assert payload["decision_score"]["formula_version"] == "dispatch-score-v1"
    assert len(payload["candidate_summaries"]) == 3
    assert sum(item["selected"] for item in payload["candidate_summaries"]) == 1
    assert payload["ai_orchestration"]["requested"] is True
    assert payload["ai_orchestration"]["used"] is False
    assert payload["ai_orchestration"]["fallback_reason"]
    cold_order_ids = {"O-001", "O-003", "O-005", "O-010"}
    for route in payload["routes"]:
        route_orders = {stop["order_id"] for stop in route["stops"]}
        if route_orders & cold_order_ids:
            assert route["cold_chain"] is True

    history = client.get("/api/v1/runs?limit=10")
    assert history.status_code == 200
    assert payload["run_id"] in {item["run_id"] for item in history.json()["items"]}
    detail = client.get(f"/api/v1/runs/{payload['run_id']}")
    assert detail.status_code == 200
    assert detail.json()["metrics"] == payload["metrics"]


def test_breakdown_reoptimization_removes_vehicle_and_persists_event(client: TestClient) -> None:
    scenario = client.get("/api/v1/scenario").json()
    cold_vehicle = next(vehicle for vehicle in scenario["vehicles"] if vehicle["cold_chain"])
    cold_order = next(order for order in scenario["orders"] if order["cold_chain"])
    event_node = next(node for node in scenario["nodes"] if node["id"] == cold_order["node_id"])
    response = client.post(
        "/api/v1/events/reoptimize",
        json={
            "event": {
                "type": "vehicle_breakdown",
                "vehicle_id": cold_vehicle["id"],
                "occurred_at_minute": 90,
                "node_id": event_node["id"],
                "longitude": event_node["longitude"],
                "latitude": event_node["latitude"],
                "location_name": event_node["name"],
            },
            "lock_completed": True,
            "urgent_orders": ["O-URGENT-01"],
        },
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] in {"optimal", "feasible"}
    assert all(route["vehicle_id"] != cold_vehicle["id"] for route in payload["routes"])
    assert payload["event_location"]["node_id"] == event_node["id"]
    assert payload["event_location"]["location_name"] == event_node["name"]
    assert payload["event_location"]["source"] == "request"
    assert payload["event_location"]["coordinate_system"] == "gcj02"

    history = client.get("/api/v1/runs?limit=20").json()["items"]
    saved = next(item for item in history if item["run_id"] == payload["run_id"])
    assert saved["event_type"] == "vehicle_breakdown"
    assert saved["event_vehicle_id"] == cold_vehicle["id"]
    with SessionLocal() as session:
        event = session.scalar(
            select(EventRecord).where(
                EventRecord.optimization_run_id == payload["run_id"]
            )
        )
        assert event is not None
        assert event.detail["location"]["node_id"] == event_node["id"]
        assert event.detail["location"]["longitude"] == pytest.approx(
            payload["event_location"]["longitude"]
        )


def test_solver_finds_depot_by_kind_not_array_position() -> None:
    scenario = get_scenario()
    depot = next(node for node in scenario.nodes if node.kind == "depot")
    reordered = scenario.model_copy(
        update={
            "nodes": [
                *[node for node in scenario.nodes if node.id != depot.id],
                depot,
            ]
        }
    )
    result = solve_scenario(reordered, time_limit_seconds=1)
    assert result.status in {"optimal", "feasible"}
    assert result.routes
    assert all(route.stops[0].node_id == depot.id for route in result.routes)
    assert all(route.stops[-1].node_id == depot.id for route in result.routes)


def test_scenario_spatial_status_detects_unpublished_coordinate_change(
    client: TestClient,
) -> None:
    scenario = client.get("/api/v1/scenario").json()
    target_id = scenario["nodes"][0]["id"]
    with SessionLocal() as session:
        target = session.get(NodeRecord, target_id)
        assert target is not None
        original_longitude = target.longitude
        target.longitude = float(target.longitude) + 0.01
        session.commit()
    try:
        response = client.get("/api/v1/scenario/spatial-status")
        assert response.status_code == 200
        payload = response.json()
        assert payload["aligned"] is False
        assert target_id in payload["changed_node_ids"]
        assert "重新发布" in payload["message"]
    finally:
        with SessionLocal() as session:
            target = session.get(NodeRecord, target_id)
            assert target is not None
            target.longitude = original_longitude
            session.commit()


def test_run_history_validation_and_missing_detail(client: TestClient) -> None:
    assert client.get("/api/v1/runs?limit=0").status_code == 422
    assert client.get("/api/v1/runs/not-found").status_code == 404


def test_latest_run_is_shared_across_roles(client: TestClient) -> None:
    optimized = client.post(
        "/api/v1/optimize",
        json={"time_limit_seconds": 1, "ai_assisted": False},
    )
    assert optimized.status_code == 200
    scenario_id = client.get("/api/v1/scenario").json()["id"]
    admin_latest = client.get(
        f"/api/v1/runs/latest?scenario_id={scenario_id}"
    )
    driver_login = client.post(
        "/api/v1/auth/login",
        json={"username": "driver", "password": "DriverDemo2026!"},
    )
    driver_headers = {
        "Authorization": f"Bearer {driver_login.json()['access_token']}"
    }
    driver_latest = client.get(
        f"/api/v1/runs/latest?scenario_id={scenario_id}",
        headers=driver_headers,
    )
    assert admin_latest.status_code == 200
    assert driver_latest.status_code == 200
    assert driver_latest.json()["run_id"] == admin_latest.json()["run_id"]
    assert driver_latest.json()["metrics"] == admin_latest.json()["metrics"]


def test_reproducible_algorithm_experiment_and_report(client: TestClient) -> None:
    assert client.get(
        "/api/v1/experiments",
        headers={"Authorization": ""},
    ).status_code == 401
    login = client.post(
        "/api/v1/auth/login",
        json={"username": "admin", "password": "QizhiDemo2026!"},
    )
    headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
    created = client.post(
        "/api/v1/experiments",
        headers=headers,
        json={
            "name": "自动化三压力对照实验",
            "time_limit_seconds": 1,
        },
    )
    assert created.status_code == 200
    payload = created.json()
    assert payload["status"] == "completed"
    assert payload["experiment_code"].startswith("EXP-")
    assert len(payload["scenario_fingerprint"]) == 64
    assert [item["case_key"] for item in payload["results"]] == [
        "standard",
        "peak",
        "cold_chain",
    ]
    assert payload["summary"]["completed_cases"] == 3
    assert payload["summary"]["average_distance_reduction_pct"] > 0
    assert isinstance(payload["summary"]["total_cost_saving_cny"], (int, float))
    assert all(item["service_rate"] == 100 for item in payload["results"])
    assert all(item["unassigned_orders"] == 0 for item in payload["results"])

    listed = client.get("/api/v1/experiments?limit=5", headers=headers)
    assert listed.status_code == 200
    assert payload["id"] in {item["id"] for item in listed.json()["items"]}
    detail = client.get(f"/api/v1/experiments/{payload['id']}", headers=headers)
    assert detail.status_code == 200
    assert detail.json()["scenario_fingerprint"] == payload["scenario_fingerprint"]
    report = client.get(
        f"/api/v1/experiments/{payload['id']}/report",
        headers=headers,
    )
    assert report.status_code == 200
    assert "text/html" in report.headers["content-type"]
    assert "祁智运算法对照实验报告" in report.text
    assert "轮询启发式" in report.text


def test_dispatch_approval_driver_execution_and_report(client: TestClient) -> None:
    admin_login = client.post(
        "/api/v1/auth/login",
        json={"username": "admin", "password": "QizhiDemo2026!"},
    )
    assert admin_login.status_code == 200
    admin_headers = {"Authorization": f"Bearer {admin_login.json()['access_token']}"}
    driver_login = client.post(
        "/api/v1/auth/login",
        json={"username": "driver", "password": "DriverDemo2026!"},
    )
    assert driver_login.status_code == 200
    assert driver_login.json()["role"] == "driver"
    driver_headers = {"Authorization": f"Bearer {driver_login.json()['access_token']}"}

    optimized = client.post("/api/v1/optimize", json={"time_limit_seconds": 1})
    assert optimized.status_code == 200
    run_id = optimized.json()["run_id"]
    submitted = client.post(
        "/api/v1/dispatch/plans",
        headers=admin_headers,
        json={"run_id": run_id},
    )
    assert submitted.status_code == 200
    plan = submitted.json()
    assert plan["status"] == "pending_approval"

    forbidden = client.post(
        f"/api/v1/dispatch/plans/{plan['id']}/approve",
        headers=driver_headers,
        json={"note": "越权审核"},
    )
    assert forbidden.status_code == 403

    approved = client.post(
        f"/api/v1/dispatch/plans/{plan['id']}/approve",
        headers=admin_headers,
        json={"note": "运力与时间窗复核通过"},
    )
    assert approved.status_code == 200
    assert approved.json()["status"] == "approved"
    repeated = client.post(
        f"/api/v1/dispatch/plans/{plan['id']}/approve",
        headers=admin_headers,
        json={},
    )
    assert repeated.status_code == 409

    issued = client.post(
        f"/api/v1/dispatch/plans/{plan['id']}/dispatch",
        headers=admin_headers,
    )
    assert issued.status_code == 200
    issued_plan = issued.json()
    assert issued_plan["status"] == "dispatched"
    assert len(issued_plan["assignments"]) == len(optimized.json()["routes"])

    assignments = client.get("/api/v1/driver/assignments", headers=driver_headers)
    assert assignments.status_code == 200
    plan_assignments = [
        item for item in assignments.json()["items"] if item["plan_id"] == plan["id"]
    ]
    assert len(plan_assignments) == len(optimized.json()["routes"])
    for assignment in plan_assignments:
        assignment_id = assignment["id"]
        assert client.post(
            f"/api/v1/driver/assignments/{assignment_id}/accept",
            headers=driver_headers,
            json={},
        ).status_code == 200
        assert client.post(
            f"/api/v1/driver/assignments/{assignment_id}/start",
            headers=driver_headers,
            json={},
        ).status_code == 200
        completed = client.post(
            f"/api/v1/driver/assignments/{assignment_id}/complete",
            headers=driver_headers,
            json={"note": "送达并完成签收"},
        )
        assert completed.status_code == 200
    assert completed.json()["status"] == "completed"

    report = client.get(
        f"/api/v1/dispatch/plans/{plan['id']}/report",
        headers=admin_headers,
    )
    assert report.status_code == 200
    assert "text/html" in report.headers["content-type"]
    assert "祁智运调度执行报告" in report.text


def test_operations_readiness_evidence_report_and_permissions(client: TestClient) -> None:
    unauthorized = client.get(
        "/api/v1/operations/readiness",
        headers={"Authorization": ""},
    )
    assert unauthorized.status_code == 401
    driver_login = client.post(
        "/api/v1/auth/login",
        json={"username": "driver", "password": "DriverDemo2026!"},
    )
    driver_headers = {
        "Authorization": f"Bearer {driver_login.json()['access_token']}"
    }
    assert client.get(
        "/api/v1/operations/readiness",
        headers=driver_headers,
    ).status_code == 403

    readiness = client.get("/api/v1/operations/readiness")
    assert readiness.status_code == 200
    payload = readiness.json()
    assert 0 <= payload["overall_score"] <= 100
    assert {item["key"] for item in payload["dimensions"]} == {
        "data",
        "algorithm",
        "execution",
        "system",
    }
    assert len(payload["script"]) == 5
    assert payload["total_duration_seconds"] == 390
    assert len(payload["kpis"]) == 4
    assert len(payload["evidence"]) >= 6
    assert any(item["evidence_id"] for item in payload["evidence"])

    report = client.get("/api/v1/operations/report")
    assert report.status_code == 200
    assert "text/html" in report.headers["content-type"]
    assert "attachment;" in report.headers["content-disposition"]
    assert "祁智运运营成果证据报告" in report.text
    assert all(word not in report.text for word in ("竞赛", "答辩", "演示", "比赛"))
    assert "API Key" in report.text
    assert "encrypted_api_key" not in report.text
    assert "QizhiDemo2026!" not in report.text


def test_map_configuration_encryption_runtime_and_permissions(client: TestClient) -> None:
    name = f"自动化地图连接-{uuid4().hex[:8]}"
    web_key = f"web-{uuid4().hex}"
    security_code = f"security-{uuid4().hex}"
    created = client.post(
        "/api/v1/map/configs",
        json={
            "name": name,
            "provider": "amap",
            "web_key": web_key,
            "security_code": security_code,
            "map_style": "whitesmoke",
            "enabled": True,
        },
    )
    assert created.status_code == 201
    config_id = created.json()["id"]
    assert created.json()["web_key_masked"] != web_key
    assert created.json()["security_code_masked"] != security_code

    configs = client.get("/api/v1/map/configs")
    assert configs.status_code == 200
    item = next(item for item in configs.json()["items"] if item["id"] == config_id)
    assert "encrypted_web_key" not in item
    assert "encrypted_security_code" not in item

    runtime = client.get("/api/v1/map/runtime")
    assert runtime.status_code == 200
    assert runtime.json()["provider"] == "amap"
    assert runtime.json()["web_key"] == web_key
    assert "security_code" not in runtime.json()
    assert runtime.json()["service_host"] == "/_AMapService"

    driver_login = client.post(
        "/api/v1/auth/login",
        json={"username": "driver", "password": "DriverDemo2026!"},
    )
    driver_headers = {
        "Authorization": f"Bearer {driver_login.json()['access_token']}"
    }
    assert client.get("/api/v1/map/configs", headers=driver_headers).status_code == 403
    driver_runtime = client.get("/api/v1/map/runtime", headers=driver_headers)
    assert driver_runtime.status_code == 200
    assert driver_runtime.json()["provider"] == "amap"
    assert "security_code" not in driver_runtime.json()

    disabled = client.post(f"/api/v1/map/configs/{config_id}/disable")
    assert disabled.status_code == 200
    assert disabled.json()["enabled"] is False
    assert client.get("/api/v1/map/runtime").json()["mode"] == "offline"

    with SessionLocal() as session:
        record = session.get(MapProviderConfigRecord, config_id)
        assert record is not None
        assert web_key not in record.encrypted_web_key
        assert security_code not in record.encrypted_security_code
        session.delete(record)
        session.commit()


def test_scenario_template_import_validation_and_activation(client: TestClient) -> None:
    login = client.post(
        "/api/v1/auth/login",
        json={"username": "admin", "password": "QizhiDemo2026!"},
    ).json()
    headers = {"Authorization": f"Bearer {login['access_token']}"}

    template = client.get("/api/v1/scenarios/template", headers=headers)
    assert template.status_code == 200
    assert template.content.startswith(b"PK")

    imported = client.post(
        "/api/v1/scenarios/import",
        headers=headers,
        files={
            "file": (
                "祁东县场景.xlsx",
                template.content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert imported.status_code == 200
    imported_scenario = imported.json()["scenario"]
    assert imported_scenario["status"] == "draft"
    assert imported_scenario["node_count"] == 5
    assert imported_scenario["order_count"] == 4
    assert imported_scenario["vehicle_count"] == 3

    activated = client.post(
        f"/api/v1/scenarios/{imported_scenario['id']}/activate",
        headers=headers,
    )
    assert activated.status_code == 200
    assert activated.json()["status"] == "active"
    assert client.get("/api/v1/scenario").json()["id"] == imported_scenario["version_key"]

    workbook = load_workbook(BytesIO(template.content))
    workbook["订单"]["B2"] = "NODE-NOT-FOUND"
    invalid_file = BytesIO()
    workbook.save(invalid_file)
    invalid = client.post(
        "/api/v1/scenarios/import",
        headers=headers,
        files={
            "file": (
                "错误场景.xlsx",
                invalid_file.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert invalid.status_code == 422
    assert "不存在的节点" in invalid.text


def test_scenario_governance_review_compare_and_rollback(client: TestClient) -> None:
    current_versions = client.get("/api/v1/scenarios").json()
    active = next(item for item in current_versions["items"] if item["status"] == "active")
    template = client.get("/api/v1/scenarios/template")
    workbook = load_workbook(BytesIO(template.content))
    workbook["订单"]["C2"] = int(workbook["订单"]["C2"].value) + 120 + active["id"]
    changed_file = BytesIO()
    workbook.save(changed_file)

    imported = client.post(
        "/api/v1/scenarios/import",
        files={
            "file": (
                "待审核场景.xlsx",
                changed_file.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert imported.status_code == 200
    target = imported.json()["scenario"]
    assert target["status"] == "draft"
    assert target["fingerprint"]

    submitted = client.post(
        f"/api/v1/scenarios/{target['id']}/submit",
        json={"note": "完成订单与运力约束校验，提交审核"},
    )
    assert submitted.status_code == 200
    assert submitted.json()["status"] == "pending_review"

    rejected = client.post(
        f"/api/v1/scenarios/{target['id']}/reject",
        json={"note": "请补充峰值订单变化说明"},
    )
    assert rejected.status_code == 200
    assert rejected.json()["status"] == "rejected"
    resubmitted = client.post(
        f"/api/v1/scenarios/{target['id']}/submit",
        json={"note": "已补充峰值货量变化说明并再次提交"},
    )
    assert resubmitted.status_code == 200

    compared = client.get(
        f"/api/v1/scenarios/compare?base_id={active['id']}&target_id={target['id']}"
    )
    assert compared.status_code == 200
    assert compared.json()["orders"]["changed"] >= 1
    assert compared.json()["deltas"]["total_weight_kg"] != 0

    approved = client.post(
        f"/api/v1/scenarios/{target['id']}/approve",
        json={"note": "审核通过，发布为生产调度场景"},
    )
    assert approved.status_code == 200
    assert approved.json()["status"] == "active"

    detail = client.get(f"/api/v1/scenarios/{target['id']}")
    assert detail.status_code == 200
    assert len(detail.json()["transitions"]) >= 5
    assert detail.json()["reviewed_by"]

    rollback = client.post(
        f"/api/v1/scenarios/{active['id']}/rollback",
        json={"note": "自动化测试结束，恢复原生产版本"},
    )
    assert rollback.status_code == 200
    assert rollback.json()["status"] == "active"
    final_versions = client.get("/api/v1/scenarios").json()
    assert final_versions["summary"]["active"] == 1
    assert final_versions["total"] >= len(final_versions["items"])


def test_ai_config_is_admin_only_encrypted_and_masked(
    client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    assert client.get(
        "/api/v1/ai/config",
        headers={"Authorization": ""},
    ).status_code == 401
    login = client.post(
        "/api/v1/auth/login",
        json={"username": "admin", "password": "QizhiDemo2026!"},
    ).json()
    headers = {"Authorization": f"Bearer {login['access_token']}"}
    raw_key = "sk-test-super-secret-value"
    mismatch = client.put(
        "/api/v1/ai/config",
        headers=headers,
        json={
            "provider": "deepseek",
            "base_url": "https://api.openai.com/v1",
            "model_name": "deepseek-v4-flash",
            "api_key": raw_key,
            "enabled": True,
            "temperature": 0.2,
        },
    )
    assert mismatch.status_code == 422
    assert "服务类型与接口地址不匹配" in mismatch.text

    saved = client.put(
        "/api/v1/ai/config",
        headers=headers,
        json={
            "provider": "openai-compatible",
            "base_url": "https://example.invalid/v1",
            "model_name": "demo-model",
            "api_key": raw_key,
            "enabled": True,
            "temperature": 0.2,
        },
    )
    assert saved.status_code == 200
    assert saved.json()["has_api_key"] is True
    assert raw_key not in saved.text
    assert "••••••" in saved.json()["api_key_masked"]

    with SessionLocal() as session:
        record = session.scalar(select(AIProviderConfigRecord))
        assert record is not None
        assert record.encrypted_api_key != raw_key
        assert raw_key not in record.encrypted_api_key

    monkeypatch.setattr("app.main.test_ai_connection", lambda _: ("连接成功", 18))
    tested = client.post("/api/v1/ai/config/test", headers=headers)
    assert tested.status_code == 200
    assert tested.json() == {"success": True, "message": "连接成功", "latency_ms": 18}

    run_id = client.get("/api/v1/runs?limit=1").json()["items"][0]["run_id"]
    monkeypatch.setattr("app.main.explain_optimization", lambda *_: ("这是可核验的调度解释。", 12))
    explanation = client.post(f"/api/v1/ai/explain/{run_id}", headers=headers)
    assert explanation.status_code == 200
    assert explanation.json()["explanation"] == "这是可核验的调度解释。"

    disabled_payload = {
        "provider": "openai-compatible",
        "base_url": "https://example.invalid/v1",
        "model_name": "demo-model",
        "api_key": None,
        "enabled": False,
        "temperature": 0.2,
    }
    disabled = client.put("/api/v1/ai/config", headers=headers, json=disabled_payload)
    assert disabled.status_code == 200
    assert disabled.json()["enabled"] is False
    assert disabled.json()["has_api_key"] is True

    visible_while_disabled = client.get("/api/v1/ai/config", headers=headers)
    assert visible_while_disabled.status_code == 200
    assert visible_while_disabled.json()["enabled"] is False

    enabled_again = client.put(
        "/api/v1/ai/config",
        headers=headers,
        json={**disabled_payload, "enabled": True},
    )
    assert enabled_again.status_code == 200
    assert enabled_again.json()["enabled"] is True
    assert enabled_again.json()["has_api_key"] is True


def test_managed_ai_connections_lifecycle(
    client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    created = client.post(
        "/api/v1/ai/configs",
        json={
            "name": "备用推理服务",
            "description": "用于企业配置管理验收",
            "provider": "openai-compatible",
            "base_url": "https://backup.example.invalid/v1",
            "model_name": "backup-model",
            "api_key": "sk-backup-secret-value",
            "enabled": False,
            "temperature": 0.3,
        },
    )
    assert created.status_code == 201
    created_payload = created.json()
    assert created_payload["name"] == "备用推理服务"
    assert created_payload["enabled"] is False
    assert "sk-backup-secret-value" not in created.text

    listed = client.get("/api/v1/ai/configs")
    assert listed.status_code == 200
    assert len(listed.json()["items"]) >= 2

    monkeypatch.setattr("app.main.test_ai_connection", lambda _: ("连接正常", 12))
    tested = client.post(f"/api/v1/ai/configs/{created_payload['id']}/test")
    assert tested.status_code == 200
    assert tested.json()["success"] is True

    activated = client.post(f"/api/v1/ai/configs/{created_payload['id']}/activate")
    assert activated.status_code == 200
    assert activated.json()["enabled"] is True
    enabled = [
        item for item in client.get("/api/v1/ai/configs").json()["items"]
        if item["enabled"]
    ]
    assert [item["id"] for item in enabled] == [created_payload["id"]]

    disabled = client.post(f"/api/v1/ai/configs/{created_payload['id']}/disable")
    assert disabled.status_code == 200
    assert disabled.json()["enabled"] is False


def test_enterprise_user_governance_and_audit_log(client: TestClient) -> None:
    assert client.get(
        "/api/v1/users",
        headers={"Authorization": ""},
    ).status_code == 401

    suffix = uuid4().hex[:8]
    username = f"ops_{suffix}"
    weak = client.post(
        "/api/v1/users",
        json={
            "username": username,
            "display_name": "运营测试账号",
            "role": "dispatcher",
            "password": "weak-password",
            "is_active": True,
        },
    )
    assert weak.status_code == 422

    created = client.post(
        "/api/v1/users",
        json={
            "username": username,
            "display_name": "运营测试账号",
            "role": "dispatcher",
            "password": "StrongPassword2026!",
            "is_active": True,
        },
    )
    assert created.status_code == 201
    user = created.json()
    assert user["username"] == username
    assert user["role"] == "dispatcher"
    assert user["is_active"] is True

    filtered = client.get(f"/api/v1/users?search={username}&role=dispatcher")
    assert filtered.status_code == 200
    assert filtered.json()["total"] == 1

    updated = client.put(
        f"/api/v1/users/{user['id']}",
        json={
            "display_name": "只读验收账号",
            "role": "viewer",
            "is_active": False,
            "unlock": True,
        },
    )
    assert updated.status_code == 200
    assert updated.json()["role"] == "viewer"
    assert updated.json()["is_active"] is False

    reset = client.post(
        f"/api/v1/users/{user['id']}/reset-password",
        json={"password": "RotatedPassword2026!"},
    )
    assert reset.status_code == 204

    admin = client.get("/api/v1/users?search=admin").json()["items"][0]
    self_demotion = client.put(
        f"/api/v1/users/{admin['id']}",
        json={
            "display_name": admin["display_name"],
            "role": "viewer",
            "is_active": True,
            "unlock": False,
        },
    )
    assert self_demotion.status_code == 409

    audit = client.get(f"/api/v1/audit-logs?category=user_management&search={user['id']}")
    assert audit.status_code == 200
    assert audit.json()["total"] >= 3
    serialized = audit.text.lower()
    assert "strongpassword2026" not in serialized
    assert "rotatedpassword2026" not in serialized

    exported = client.get("/api/v1/audit-logs/export?category=user_management")
    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith("text/csv")
    assert exported.content.startswith(b"\xef\xbb\xbf")


def test_master_data_center_crud_validation_and_audit(client: TestClient) -> None:
    assert client.get(
        "/api/v1/master-data/summary",
        headers={"Authorization": ""},
    ).status_code == 401

    suffix = uuid4().hex[:8]
    node_id = f"N-QA-{suffix}"
    vehicle_id = f"V-QA-{suffix}"
    order_id = f"O-QA-{suffix}"

    node = client.post(
        "/api/v1/master-data/nodes",
        json={
            "id": node_id,
            "name": "基础数据验收网点",
            "longitude": 112.11,
            "latitude": 26.79,
            "kind": "town",
            "address": "衡阳市祁东县验收路 1 号",
            "contact_name": "验收联系人",
            "contact_phone": "13800000000",
            "service_radius_km": 18.5,
            "is_active": True,
        },
    )
    assert node.status_code == 201
    assert node.json()["is_active"] is True

    vehicle = client.post(
        "/api/v1/master-data/vehicles",
        json={
            "id": vehicle_id,
            "plate": f"湘DQA{suffix[:4].upper()}",
            "capacity_kg": 1800,
            "cold_chain": True,
            "vehicle_type": "refrigerated",
            "driver_name": "车辆验收员",
            "driver_phone": "13900000000",
            "status": "available",
            "cost_per_km": 3.8,
            "shift_start": 420,
            "shift_end": 1080,
        },
    )
    assert vehicle.status_code == 201
    assert vehicle.json()["available"] is True

    order_payload = {
        "node_id": node_id,
        "cargo_name": "黄花菜",
        "customer_name": "验收合作社",
        "customer_phone": "13700000000",
        "delivery_address": "祁东县验收镇",
        "weight_kg": 650,
        "window_start": 480,
        "window_end": 720,
        "service_minutes": 12,
        "cold_chain": True,
        "priority": "urgent",
        "status": "pending",
    }
    order = client.post(
        "/api/v1/master-data/orders",
        json={"id": order_id, **order_payload},
    )
    assert order.status_code == 201
    assert order.json()["node_name"] == "基础数据验收网点"

    filtered = client.get(
        f"/api/v1/master-data/orders?search={order_id}&status=pending&priority=urgent&cold_chain=true"
    )
    assert filtered.status_code == 200
    assert filtered.json()["total"] == 1

    maintenance = client.put(
        f"/api/v1/master-data/vehicles/{vehicle_id}",
        json={
            **{key: value for key, value in vehicle.json().items() if key in {
                "plate", "capacity_kg", "cold_chain", "vehicle_type", "driver_name",
                "driver_phone", "cost_per_km", "shift_start", "shift_end",
            }},
            "status": "maintenance",
        },
    )
    assert maintenance.status_code == 200
    assert maintenance.json()["available"] is False

    inactive = client.put(
        f"/api/v1/master-data/nodes/{node_id}",
        json={
            **{key: value for key, value in node.json().items() if key in {
                "name", "longitude", "latitude", "kind", "address", "contact_name",
                "contact_phone", "service_radius_km",
            }},
            "is_active": False,
        },
    )
    assert inactive.status_code == 200

    rejected_order = client.post(
        "/api/v1/master-data/orders",
        json={"id": f"O-BLOCK-{suffix}", **order_payload},
    )
    assert rejected_order.status_code == 422
    assert "已停用" in rejected_order.text

    protected_node = client.delete(f"/api/v1/master-data/nodes/{node_id}")
    assert protected_node.status_code == 409

    deleted_order = client.delete(f"/api/v1/master-data/orders/{order_id}")
    assert deleted_order.status_code == 204
    assert client.delete(f"/api/v1/master-data/vehicles/{vehicle_id}").status_code == 204
    assert client.delete(f"/api/v1/master-data/nodes/{node_id}").status_code == 204

    summary = client.get("/api/v1/master-data/summary")
    assert summary.status_code == 200
    assert summary.json()["order_total"] >= 1

    audit = client.get(f"/api/v1/audit-logs?category=master_data&search={node_id}")
    assert audit.status_code == 200
    assert audit.json()["total"] >= 3


def test_node_coordinate_calibration_and_revision_history(client: TestClient) -> None:
    suffix = uuid4().hex[:8]
    node_id = f"N-SP-{suffix}"
    created = client.post(
        "/api/v1/master-data/nodes",
        json={
            "id": node_id,
            "name": "空间校准测试网点",
            "longitude": 112.08,
            "latitude": 26.78,
            "kind": "town",
            "address": "衡阳市祁东县测试地址",
            "service_radius_km": 15,
            "is_active": True,
        },
    )
    assert created.status_code == 201
    assert created.json()["coordinate_status"] == "unverified"
    assert created.json()["coordinate_system"] == "gcj02"
    assert created.json()["wgs84_longitude"] is not None

    calibrated = client.put(
        f"/api/v1/master-data/nodes/{node_id}/coordinate",
        json={
            "longitude": 112.091234,
            "latitude": 26.791234,
            "address": "湖南省衡阳市祁东县校准路 8 号",
            "coordinate_source": "manual",
            "coordinate_accuracy_m": 12,
            "coordinate_status": "verified",
            "note": "现场定位核验",
        },
    )
    assert calibrated.status_code == 200
    payload = calibrated.json()
    assert payload["coordinate_status"] == "verified"
    assert payload["coordinate_accuracy_m"] == 12
    assert payload["longitude"] == pytest.approx(112.091234)
    assert payload["wgs84_longitude"] != pytest.approx(payload["longitude"])

    history = client.get(
        f"/api/v1/master-data/nodes/{node_id}/coordinate-revisions"
    )
    assert history.status_code == 200
    assert history.json()["items"][0]["note"] == "现场定位核验"
    assert history.json()["items"][0]["actor_username"] == "admin"

    quality = client.get("/api/v1/master-data/spatial-quality")
    assert quality.status_code == 200
    assert quality.json()["display_coordinate_system"] == "GCJ-02"
    assert quality.json()["storage_coordinate_system"] == "WGS-84"
    assert quality.json()["verified"] >= 1

    audit = client.get(
        f"/api/v1/audit-logs?category=master_data&search={node_id}"
    )
    assert audit.status_code == 200
    assert any(
        item["action"] == "calibrate_node_coordinate"
        for item in audit.json()["items"]
    )
    assert client.delete(f"/api/v1/master-data/nodes/{node_id}").status_code == 204


def test_order_import_quality_rules_and_master_data_publish(client: TestClient) -> None:
    template = client.get("/api/v1/master-data/orders/import/template")
    assert template.status_code == 200
    workbook = load_workbook(BytesIO(template.content), read_only=True)
    assert workbook.active.title == "订单导入"
    assert workbook.active["A1"].value == "订单编号"

    suffix = uuid4().hex[:8]
    active_nodes = client.get("/api/v1/master-data/nodes?active=true&limit=100")
    assert active_nodes.status_code == 200
    node_id = active_nodes.json()["items"][0]["id"]
    header = (
        "订单编号,货物名称,客户名称,联系电话,服务网点编号,配送地址,"
        "重量(kg),时间窗开始,时间窗结束,服务时长(分钟),冷链要求,优先级\r\n"
    )
    bad_csv = (
        header
        + f"O-IMP-BAD-{suffix},黄花菜,错误合作社,13800000000,N-NOT-EXIST,"
        "祁东县测试地址,500,08:00,12:00,8,否,普通\r\n"
    ).encode("utf-8-sig")
    bad = client.post(
        "/api/v1/master-data/orders/import/preview",
        files={"file": (f"bad-{suffix}.csv", bad_csv, "text/csv")},
    )
    assert bad.status_code == 201
    assert bad.json()["status"] == "rejected"
    assert bad.json()["error_rows"] == 1
    assert bad.json()["errors"][0]["field"] == "服务网点编号"

    receipt = client.get(f"/api/v1/master-data/imports/{bad.json()['id']}/errors")
    assert receipt.status_code == 200
    assert receipt.content.startswith(b"\xef\xbb\xbf")
    assert "网点不存在或未启用" in receipt.content.decode("utf-8-sig")

    order_id = f"O-IMP-{suffix}"
    good_csv = (
        header
        + f"{order_id},黄花菜,批量导入合作社,13800000000,{node_id},"
        "祁东县测试地址,680,08:00,12:00,8,否,紧急\r\n"
    ).encode("utf-8-sig")
    good = client.post(
        "/api/v1/master-data/orders/import/preview",
        files={"file": (f"good-{suffix}.csv", good_csv, "text/csv")},
    )
    assert good.status_code == 201
    assert good.json()["status"] == "validated"
    assert good.json()["valid_rows"] == 1

    committed = client.post(
        f"/api/v1/master-data/imports/{good.json()['id']}/commit"
    )
    assert committed.status_code == 200
    assert committed.json()["status"] == "committed"
    imported = client.get(f"/api/v1/master-data/orders?search={order_id}")
    assert imported.status_code == 200
    assert imported.json()["total"] == 1

    quality = client.get("/api/v1/master-data/quality-report")
    assert quality.status_code == 200
    assert len(quality.json()["items"]) == 5
    assert {"blocked_count", "warning_count", "publish_ready"} <= quality.json().keys()
    first_rule = quality.json()["items"][0]
    updated_rule = client.put(
        f"/api/v1/master-data/quality-rules/{first_rule['id']}",
        json={
            "enabled": first_rule["enabled"],
            "severity": first_rule["severity"],
        },
    )
    assert updated_rule.status_code == 200
    assert updated_rule.json()["rule_code"] == first_rule["rule_code"]

    published_id = None
    if quality.json()["publish_ready"]:
        published = client.post(
            "/api/v1/master-data/publish",
            json={"name": f"批量导入验收场景 {suffix}", "activate": False},
        )
        assert published.status_code == 201
        assert published.json()["scenario"]["status"] == "draft"
        assert published.json()["scenario"]["scenario_code"] == "master-data"
        with SessionLocal() as session:
            snapshot = session.get(
                ScenarioVersionRecord, published.json()["scenario"]["id"]
            ).payload
            assert snapshot["spatial_snapshot_id"].startswith("spatial-")
            assert snapshot["spatial_snapshot_at"]
            assert snapshot["nodes"][0]["kind"] == "depot"
        published_id = published.json()["scenario"]["id"]

    assert client.delete(f"/api/v1/master-data/orders/{order_id}").status_code == 204
    with SessionLocal() as session:
        for batch_id in (bad.json()["id"], good.json()["id"]):
            batch = session.get(DataImportBatchRecord, batch_id)
            if batch:
                session.delete(batch)
        if published_id:
            scenario = session.get(ScenarioVersionRecord, published_id)
            if scenario:
                session.delete(scenario)
        stale_order = session.get(OrderRecord, order_id)
        if stale_order:
            session.delete(stale_order)
        session.commit()
