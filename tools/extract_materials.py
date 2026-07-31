from pathlib import Path
import json
import sys
import zipfile

from docx import Document


def extract_docx(path: Path, out_dir: Path) -> dict:
    doc = Document(path)
    blocks = []
    for i, p in enumerate(doc.paragraphs, 1):
        text = p.text.strip()
        if text:
            blocks.append({"type": "paragraph", "index": i, "style": p.style.name, "text": text})
    for ti, table in enumerate(doc.tables, 1):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        blocks.append({"type": "table", "index": ti, "rows": rows})

    media_dir = out_dir / f"{path.stem}_media"
    media = []
    with zipfile.ZipFile(path) as zf:
        for name in zf.namelist():
            if name.startswith("word/media/") and not name.endswith("/"):
                media_dir.mkdir(parents=True, exist_ok=True)
                target = media_dir / Path(name).name
                target.write_bytes(zf.read(name))
                media.append(str(target))

    result = {
        "source": str(path),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "sections": len(doc.sections),
        "blocks": blocks,
        "media": media,
    }
    output = out_dir / f"{path.stem}.json"
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def main() -> None:
    out_dir = Path(sys.argv[1])
    out_dir.mkdir(parents=True, exist_ok=True)
    summary = []
    for value in sys.argv[2:]:
        result = extract_docx(Path(value), out_dir)
        summary.append({
            "source": result["source"],
            "paragraph_count": result["paragraph_count"],
            "table_count": result["table_count"],
            "sections": result["sections"],
            "media_count": len(result["media"]),
        })
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
