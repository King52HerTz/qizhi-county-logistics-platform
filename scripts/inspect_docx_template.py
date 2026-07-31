from __future__ import annotations

import json
import sys
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def value(value):
    return None if value is None else {
        "emu": int(value),
        "cm": round(value.cm, 3),
        "pt": round(value.pt, 2),
    }


source = Path(sys.argv[1])
document = Document(source)
styles = {}
for name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Caption"]:
    try:
        style = document.styles[name]
    except KeyError:
        styles[name] = None
        continue
    pf = style.paragraph_format
    styles[name] = {
        "font": style.font.name,
        "size_pt": style.font.size.pt if style.font.size else None,
        "bold": style.font.bold,
        "italic": style.font.italic,
        "color": str(style.font.color.rgb) if style.font.color and style.font.color.rgb else None,
        "alignment": str(pf.alignment),
        "left_indent": value(pf.left_indent),
        "first_line_indent": value(pf.first_line_indent),
        "space_before": value(pf.space_before),
        "space_after": value(pf.space_after),
        "line_spacing": str(pf.line_spacing),
    }

sections = []
for section in document.sections:
    sections.append({
        "width": value(section.page_width),
        "height": value(section.page_height),
        "orientation": str(section.orientation),
        "top_margin": value(section.top_margin),
        "bottom_margin": value(section.bottom_margin),
        "left_margin": value(section.left_margin),
        "right_margin": value(section.right_margin),
        "header_distance": value(section.header_distance),
        "footer_distance": value(section.footer_distance),
    })

style_usage = Counter(p.style.name for p in document.paragraphs)
with zipfile.ZipFile(source) as archive:
    package = {
        info.filename: {"size": info.file_size, "crc": info.CRC}
        for info in archive.infolist()
    }

result = {
    "paragraphs": len(document.paragraphs),
    "tables": len(document.tables),
    "inline_shapes": len(document.inline_shapes),
    "sections": sections,
    "styles": styles,
    "style_usage": style_usage,
    "package": package,
}
print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
