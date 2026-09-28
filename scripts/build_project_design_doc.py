from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path("E:/wuliu_project")
# The previous document is kept as a reference artifact.  This builder creates
# a separate, formal edition so the original file is never overwritten.
OUTPUT = ROOT / "祁智运县域物流智能决策平台设计与实现-正式稿.docx"
ASSETS = ROOT / "docs/assets/project-document"

GREEN = "#0D5B3D"
GREEN_2 = "#178052"
PALE = "#EAF4EE"
AMBER = "#F3B63E"
SLATE = "#30473B"
GRAY = "#66776E"
LINE = "#C9D8CF"
WHITE = "#FFFFFF"


def font(size: int, bold: bool = False, color: str = SLATE):
    candidates = [
        "C:/Windows/Fonts/msyh.ttc",
        "C:/Windows/Fonts/simhei.ttf",
        "C:/Windows/Fonts/simsun.ttc",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size=size, index=0)
    return ImageFont.load_default()


def arrow(draw, start, end, fill=GREEN_2, width=5):
    draw.line([start, end], fill=fill, width=width)
    x2, y2 = end
    x1, y1 = start
    dx, dy = x2 - x1, y2 - y1
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    head = 16
    points = [
        (x2, y2),
        (x2 - ux * head + px * head * 0.55, y2 - uy * head + py * head * 0.55),
        (x2 - ux * head - px * head * 0.55, y2 - uy * head - py * head * 0.55),
    ]
    draw.polygon(points, fill=fill)


def box(draw, xy, title, subtitle="", fill=WHITE, outline=LINE, title_color=SLATE):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=3)
    title_f = font(30, True, title_color)
    subtitle_f = font(19, False, GRAY)
    title_box = draw.textbbox((0, 0), title, font=title_f)
    title_w = title_box[2] - title_box[0]
    base_y = y1 + 29 if subtitle else (y1 + y2 - (title_box[3] - title_box[1])) / 2 - 4
    draw.text(((x1 + x2 - title_w) / 2, base_y), title, font=title_f, fill=title_color)
    if subtitle:
        subtitle_box = draw.textbbox((0, 0), subtitle, font=subtitle_f)
        subtitle_w = subtitle_box[2] - subtitle_box[0]
        draw.text(((x1 + x2 - subtitle_w) / 2, base_y + 48), subtitle, font=subtitle_f, fill=GRAY)


def diagram_canvas(title, subtitle):
    image = Image.new("RGB", (1600, 900), "#F7F9F7")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((45, 40, 1555, 860), radius=30, fill=WHITE, outline="#DEE8E1", width=3)
    draw.text((95, 82), title, font=font(42, True, GREEN), fill=GREEN)
    draw.text((98, 140), subtitle, font=font(22, False, GRAY), fill=GRAY)
    draw.line((95, 184, 1505, 184), fill="#DDE8E1", width=3)
    return image, draw


def create_diagrams():
    # Business decision loop.
    image, draw = diagram_canvas("县域物流智能决策闭环", "从可信数据到可执行任务，再由动态事件反馈驱动重调度")
    items = [
        ("订单与运力", "数据校验 / 场景版本"),
        ("需求预测", "统计基线 / 风险识别"),
        ("策略编排", "AI 可选 / 白名单"),
        ("约束求解", "OR-Tools / 硬约束"),
        ("透明评分", "里程 / 成本 / 服务"),
        ("审批执行", "下发 / 接单 / 完成"),
    ]
    positions = [(90 + i * 245, 280, 290 + i * 245, 425) for i in range(6)]
    for idx, (title, sub) in enumerate(items):
        fill = PALE if idx not in (2, 3) else ("#FFF5DD" if idx == 2 else "#DDF1E6")
        box(draw, positions[idx], title, sub, fill=fill)
        if idx < len(items) - 1:
            arrow(draw, (positions[idx][2] + 8, 352), (positions[idx + 1][0] - 8, 352))
    draw.rounded_rectangle((250, 570, 1350, 740), radius=28, fill="#F1F6F3", outline=LINE, width=3)
    draw.text((320, 612), "动态事件：车辆故障 / 紧急订单 / 路网变化", font=font(31, True, SLATE), fill=SLATE)
    draw.text((385, 674), "冻结已执行任务，仅重构剩余任务求解实例", font=font(24, False, GREEN_2), fill=GREEN_2)
    arrow(draw, (1275, 425), (1275, 570))
    arrow(draw, (250, 655), (125, 425))
    image.save(ASSETS / "10-business-loop.png")

    # Technical architecture.
    image, draw = diagram_canvas("系统技术架构", "容器化部署、前后端分离、同步 API 与异步求解协同")
    layers = [
        ("接入层", ["浏览器", "Nginx", "高德 JSAPI / 离线地图"]),
        ("应用层", ["Vue 3 + TypeScript", "运营总览", "调度 / 事件 / 执行 / 设置"]),
        ("服务层", ["FastAPI", "Celery Worker", "JWT / RBAC / 审计"]),
        ("智能层", ["OR-Tools", "OSRM 路网矩阵", "AI 策略建议与复核"]),
        ("数据层", ["PostgreSQL + PostGIS", "Redis", "Alembic 迁移"]),
    ]
    y = 230
    for idx, (layer, elements) in enumerate(layers):
        draw.rounded_rectangle((110, y, 1490, y + 105), radius=16, fill=PALE if idx % 2 == 0 else "#F8F4E8", outline=LINE, width=2)
        draw.rounded_rectangle((130, y + 17, 310, y + 88), radius=13, fill=GREEN if idx < 4 else SLATE)
        label_box = draw.textbbox((0, 0), layer, font=font(28, True, WHITE))
        draw.text((220 - (label_box[2] - label_box[0]) / 2, y + 34), layer, font=font(28, True, WHITE), fill=WHITE)
        widths = [355, 750, 1145]
        for x, element in zip(widths, elements):
            draw.text((x, y + 37), element, font=font(24, True if idx in (2, 3) else False, SLATE), fill=SLATE)
        y += 120
    image.save(ASSETS / "11-technical-architecture.png")

    # AI + deterministic boundary.
    image, draw = diagram_canvas("混合智能决策边界", "大模型参与目标理解和安全候选复核，确定性求解器守住业务硬约束")
    left = (110, 250, 500, 700)
    middle = (605, 250, 995, 700)
    right = (1100, 250, 1490, 700)
    draw.rounded_rectangle(left, radius=24, fill="#FFF5DD", outline="#E9C976", width=4)
    draw.rounded_rectangle(middle, radius=24, fill="#DDF1E6", outline="#7FB89A", width=4)
    draw.rounded_rectangle(right, radius=24, fill="#EDF2F8", outline="#AAB9C9", width=4)
    draw.text((205, 290), "AI 协同层", font=font(36, True, "#8A5B00"), fill="#8A5B00")
    draw.text((660, 290), "确定性求解层", font=font(36, True, GREEN), fill=GREEN)
    draw.text((1190, 290), "证据与治理层", font=font(36, True, "#3E5368"), fill="#3E5368")
    left_lines = ["理解业务目标", "推荐策略参数", "复核安全候选", "失败自动降级"]
    middle_lines = ["容量与时间窗", "冷链与车辆班次", "可行性校验", "OR-Tools 生成路线"]
    right_lines = ["公开评分公式", "场景与路网快照", "候选方案留痕", "审批与执行审计"]
    for lines, x, color in [(left_lines, 170, "#6B5A31"), (middle_lines, 665, SLATE), (right_lines, 1160, "#4D6071")]:
        for idx, line in enumerate(lines):
            draw.ellipse((x, 390 + idx * 64, x + 18, 408 + idx * 64), fill=AMBER if x < 500 else GREEN_2)
            draw.text((x + 34, 378 + idx * 64), line, font=font(25, False, color), fill=color)
    arrow(draw, (510, 475), (595, 475))
    arrow(draw, (1005, 475), (1090, 475))
    image.save(ASSETS / "12-ai-boundary.png")

    # Deployment.
    image, draw = diagram_canvas("Docker Compose 部署拓扑", "Nginx 统一入口，迁移先行，业务与数据服务独立持久化")
    boxes = {
        "user": (110, 345, 350, 490),
        "nginx": (455, 345, 695, 490),
        "api": (805, 245, 1060, 390),
        "worker": (805, 505, 1060, 650),
        "db": (1180, 245, 1440, 390),
        "redis": (1180, 505, 1440, 650),
    }
    box(draw, boxes["user"], "浏览器", "http://127.0.0.1", fill="#F4F6F5")
    box(draw, boxes["nginx"], "Nginx", "静态资源 / API 代理", fill=PALE)
    box(draw, boxes["api"], "FastAPI", "认证 / 业务 / 同步求解", fill="#DDF1E6")
    box(draw, boxes["worker"], "Celery Worker", "异步优化任务", fill="#FFF5DD")
    box(draw, boxes["db"], "PostgreSQL", "业务数据 / 配置 / 审计", fill="#EDF2F8")
    box(draw, boxes["redis"], "Redis", "队列 / 结果 / 缓存", fill="#F6EFEA")
    arrow(draw, (350, 418), (445, 418))
    arrow(draw, (695, 418), (795, 318))
    arrow(draw, (695, 418), (795, 578))
    arrow(draw, (1060, 318), (1170, 318))
    arrow(draw, (1060, 578), (1170, 578))
    arrow(draw, (1060, 345), (1170, 540), fill="#759284", width=4)
    draw.text((490, 710), "迁移容器在启动时执行 Alembic 升级后退出；PostgreSQL 与 Redis 使用独立数据卷", font=font(24, False, GRAY), fill=GRAY)
    image.save(ASSETS / "13-deployment-topology.png")


def clear_body(document: Document):
    body = document._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def set_east_asia(run, name="宋体"):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def configure_styles(document: Document):
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # Match the supplied reference document: body paragraphs are flush-left
    # and the visible paragraph rhythm is carried by exact line spacing.
    normal.paragraph_format.first_line_indent = None
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(20)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    for name, size, before, after in [
        ("Heading 1", 15, 0, 10),
        ("Heading 2", 14, 13, 13),
        ("Heading 3", 12, 13, 8),
    ]:
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        pf = style.paragraph_format
        # The reference uses hanging indents so that manually numbered
        # headings line up cleanly even when the number has two digits.
        if name == "Heading 1":
            pf.left_indent = Cm(0.762)
            pf.first_line_indent = Cm(-0.762)
        elif name == "Heading 2":
            pf.left_indent = Cm(1.438)
            pf.first_line_indent = Cm(-1.014)
        else:
            pf.left_indent = Cm(2.117)
            pf.first_line_indent = Cm(-1.27)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(30 if name != "Heading 3" else 26)
        pf.keep_with_next = True

    caption = document.styles["Caption"]
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(10)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Pt(0)
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(0)
    caption.paragraph_format.keep_with_previous = True

    # The supplied document uses dedicated 图/表 caption styles.  Recreate
    # those style names so captions remain consistent with the reference even
    # when the file is opened in Word or WPS.
    for style_name in ("图", "表"):
        if style_name in [style.name for style in document.styles]:
            style = document.styles[style_name]
        else:
            style = document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Times New Roman"
        style.font.size = Pt(10)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1
        style.paragraph_format.keep_with_previous = True

    # A muted utility style is used for source notes and document-control
    # metadata.  Keeping this as a real style makes the visual system stable
    # when the DOCX is opened in Word or LibreOffice.
    if "Qizhi Note" not in [style.name for style in document.styles]:
        note = document.styles.add_style("Qizhi Note", 1)
    else:
        note = document.styles["Qizhi Note"]
    note.font.name = "Times New Roman"
    note.font.size = Pt(9.5)
    note.font.color.rgb = RGBColor(102, 119, 110)
    note._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    note.paragraph_format.first_line_indent = Pt(0)
    note.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    note.paragraph_format.line_spacing = 1.08
    note.paragraph_format.space_before = Pt(2)
    note.paragraph_format.space_after = Pt(7)


def add_page_field(paragraph):
    """Insert a real Word PAGE field instead of a static page number."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_east_asia(run)
    run.font.size = Pt(9)


def configure_header_footer(document: Document):
    """Apply the understated running header/footer used by the current edition."""
    for section in document.sections:
        section.different_first_page_header_footer = True
        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.text = ""
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hp.paragraph_format.first_line_indent = Pt(0)
        hp.paragraph_format.space_after = Pt(0)
        hp.paragraph_format.line_spacing = 1
        hr = hp.add_run("祁智运  ·  县域物流智能决策平台  ·  设计与实现")
        set_east_asia(hr)
        hr.font.size = Pt(9)
        hr.font.color.rgb = RGBColor(102, 119, 110)

        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.text = ""
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.first_line_indent = Pt(0)
        fp.paragraph_format.space_before = Pt(0)
        fp.paragraph_format.space_after = Pt(0)
        fp.paragraph_format.line_spacing = 1
        left = fp.add_run("项目设计资料  ·  ")
        set_east_asia(left)
        left.font.size = Pt(9)
        left.font.color.rgb = RGBColor(139, 151, 145)
        mid = fp.add_run("第 ")
        set_east_asia(mid)
        mid.font.size = Pt(9)
        mid.font.color.rgb = RGBColor(139, 151, 145)
        add_page_field(fp)
        right = fp.add_run(" 页")
        set_east_asia(right)
        right.font.size = Pt(9)
        right.font.color.rgb = RGBColor(139, 151, 145)


def add_cover_text(document: Document, text: str, size: int, *, bold=False,
                   color=SLATE, align=WD_ALIGN_PARAGRAPH.CENTER,
                   before=0, after=0):
    paragraph = document.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1
    run = paragraph.add_run(text)
    set_east_asia(run, "黑体" if bold else "宋体")
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))
    return paragraph


def add_source_note(document: Document, text: str):
    paragraph = document.add_paragraph(text, style="Qizhi Note")
    for run in paragraph.runs:
        set_east_asia(run)
    return paragraph


def add_callout(document: Document, label: str, text: str, *, fill=PALE,
                accent=GREEN_2):
    """Add a compact, accessible callout without using a layout table."""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.left_indent = Pt(8)
    paragraph.paragraph_format.right_indent = Pt(4)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(9)
    paragraph.paragraph_format.line_spacing = 1.12
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill.replace("#", ""))
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), accent.replace("#", ""))
    borders.append(left)
    p_pr.append(borders)
    lead = paragraph.add_run(f"{label}  ")
    set_east_asia(lead, "黑体")
    lead.bold = True
    lead.font.size = Pt(11)
    lead.font.color.rgb = RGBColor.from_string(accent.replace("#", ""))
    body = paragraph.add_run(text)
    set_east_asia(body)
    body.font.size = Pt(11)
    body.font.color.rgb = RGBColor.from_string(SLATE.replace("#", ""))
    return paragraph


def add_metadata_table(document: Document, rows):
    return add_table(
        document,
        "文档控制信息",
        ["字段", "内容"],
        rows,
        [3.0, 11.4],
    )


def add_heading(document, text, level):
    paragraph = document.add_paragraph(style=f"Heading {level}")
    # The reference template links its heading styles to a multilevel list.
    # The project headings already contain the required 6.x numbering, so an
    # explicit numId=0 suppresses inherited list numbering in Word.
    p_pr = paragraph._p.get_or_add_pPr()
    inherited_num_pr = p_pr.find(qn("w:numPr"))
    if inherited_num_pr is not None:
        p_pr.remove(inherited_num_pr)
    num_pr = OxmlElement("w:numPr")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "0")
    num_pr.append(num_id)
    p_pr.append(num_pr)
    run = paragraph.add_run(text)
    set_east_asia(run, "黑体")
    run.bold = True
    return paragraph


def add_body(document, text, bold_prefix=None):
    paragraph = document.add_paragraph(style="Normal")
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        set_east_asia(run, "黑体")
        run.bold = True
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_east_asia(rest)
    else:
        run = paragraph.add_run(text)
        set_east_asia(run)
    return paragraph


def add_numbered(document, items):
    for index, item in enumerate(items, 1):
        paragraph = document.add_paragraph(style="Normal")
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.left_indent = Pt(0)
        paragraph.paragraph_format.hanging_indent = Pt(0)
        run = paragraph.add_run(f"{index}. {item}")
        set_east_asia(run)


def add_figure(document, filename, caption, width=14.5):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    # Normal uses an exact 20 pt line height in the academic template. Inline
    # pictures would therefore be clipped to a thin strip unless the figure
    # paragraph explicitly allows its line box to grow with the image.
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(0)
    inline_shape = paragraph.add_run().add_picture(str(ASSETS / filename), width=Cm(width))
    # Store meaningful alternative text in the drawing metadata so the
    # document remains understandable to screen-reader users.
    inline_shape._inline.docPr.set("descr", caption)
    inline_shape._inline.docPr.set("title", caption)
    cap = document.add_paragraph(caption, style="图")
    cap.paragraph_format.keep_with_previous = True
    return cap


def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(str(text))
    set_east_asia(run)
    run.font.size = Pt(10.5)
    run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def border(cell, edge, size=8, color="000000"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    tag = qn(f"w:{edge}")
    node = borders.find(tag)
    if node is None:
        node = OxmlElement(f"w:{edge}")
        borders.append(node)
    node.set(qn("w:val"), "single")
    node.set(qn("w:sz"), str(size))
    node.set(qn("w:color"), color)


def remove_border(cell, edge):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    node = borders.find(qn(f"w:{edge}"))
    if node is None:
        node = OxmlElement(f"w:{edge}")
        borders.append(node)
    node.set(qn("w:val"), "nil")


def add_table(document, caption, headers, rows, widths=None):
    cap = document.add_paragraph(caption, style="表")
    cap.paragraph_format.keep_with_next = True
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_pr = table._tbl.tblPr
    table_width_cm = sum(widths) if widths else 14.4
    table_width_dxa = int(table_width_cm / 2.54 * 1440)
    tbl_w = table_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        table_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(table_width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = table_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        table_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_layout = table_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        table_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_grid = table._tbl.tblGrid
    for grid_col in list(tbl_grid):
        tbl_grid.remove(grid_col)
    grid_widths = widths or [table_width_cm / len(headers)] * len(headers)
    for width in grid_widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width / 2.54 * 1440)))
        tbl_grid.append(grid_col)
    header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_tr_pr.append(repeat_header)
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=False)
        # Keep the heading row with the first data row, and repeat it if the
        # table subsequently spans pages.
        table.rows[0].cells[idx].paragraphs[0].paragraph_format.keep_with_next = True
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.LEFT if idx == len(row) - 1 and len(str(value)) > 12 else WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(cells[idx], value, align=align)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
                tc_pr = row.cells[idx]._tc.get_or_add_tcPr()
                tc_w = tc_pr.first_child_found_in("w:tcW")
                if tc_w is None:
                    tc_w = OxmlElement("w:tcW")
                    tc_pr.append(tc_w)
                tc_w.set(qn("w:w"), str(int(width / 2.54 * 1440)))
                tc_w.set(qn("w:type"), "dxa")
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_margins(cell)
            for edge in ("left", "right", "insideH", "insideV"):
                remove_border(cell, edge)
            if row_index == 0:
                border(cell, "top", 12)
                border(cell, "bottom", 6)
            elif row_index == len(table.rows) - 1:
                border(cell, "bottom", 12)
            else:
                remove_border(cell, "top")
                remove_border(cell, "bottom")
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_equation(document, text, number):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"{text}    （{number}）")
    run.font.name = "Cambria Math"
    run.font.size = Pt(12)
    run.italic = True


def renumber_figures(document):
    """Keep figure numbering sequential after inserting generated diagrams."""
    figure_number = 0
    for paragraph in document.paragraphs:
        match = re.match(r"^\s*图\s+6-\d+\s*", paragraph.text)
        if not match:
            continue
        figure_number += 1
        description = paragraph.text[match.end():]
        paragraph.text = f"图 6-{figure_number} {description}"


def build():
    ASSETS.mkdir(parents=True, exist_ok=True)
    create_diagrams()
    document = Document()
    configure_styles(document)
    section = document.sections[0]
    section.page_width = Cm(21.001)
    section.page_height = Cm(29.7)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.175)
    section.right_margin = Cm(3.175)
    section.header_distance = Cm(1.501)
    section.footer_distance = Cm(1.75)

    # Cover page.  It intentionally avoids the built-in Word Title style so
    # the document has no template residue when opened in different editors.
    add_cover_text(document, "祁智运", 18, bold=True, color=GREEN, before=26, after=12)
    add_cover_text(document, "县域物流智能决策平台", 30, bold=True, color=SLATE, after=6)
    add_cover_text(document, "设计与实现", 22, bold=True, color=GREEN_2, after=14)
    add_cover_text(document, "基于祁东县县乡村三级物流网络的智能调度、路径规划与运营闭环", 12, color=GRAY, after=24)
    rule = document.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rule.paragraph_format.first_line_indent = Pt(0)
    rule.paragraph_format.space_after = Pt(20)
    p_pr = rule._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "22")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), GREEN_2.replace("#", ""))
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    add_figure(document, "28-county-logistics-network-image2.png", "图 0-1 祁东县县域物流三级配送网络示意", 13.9)
    document.add_page_break()

    # Chapter 1 — executive view and boundaries.
    add_heading(document, "1  文档概览", 1)
    add_heading(document, "1.1 编写目的", 2)
    add_body(document, "本文件将祁智运平台整理为一份可阅读、可部署、可验收的系统设计与实现说明。它回答三个问题：平台解决什么物流问题，平台已经形成哪些业务闭环，使用者如何通过操作界面完成启动、配置、运行和恢复。文档同时把新增研究资料转化为可追溯的数据依据，避免把规划研究数据、标准场景数据和实时生产数据混为一谈。")
    add_callout(document, "先看结论", "祁智运是一套面向祁东县县域物流的智能决策平台：确定性 OR-Tools 求解负责路线可行性，OSRM 提供道路矩阵，高德地图负责可视化，AI 服务只在配置并启用后参与目标理解与候选复核。即使外部 AI 或地图服务暂时不可用，核心调度流程仍可依靠降级能力完成。")
    add_heading(document, "1.2 平台建设边界", 2)
    add_table(document, "表 1-1 平台建设范围与边界", ["范围", "已实现能力", "明确不作的承诺"], [
        ["业务对象", "订单、车辆、节点、方案、事件、执行任务、实验与审计", "不等同于企业 TMS/WMS 的全部生产模块"],
        ["优化能力", "容量、时间窗、冷链、班次、车辆可用性与动态重调度", "不宣称对所有现实约束都达到全局最优"],
        ["空间能力", "县域边界、校准节点、OSRM 距离/时间矩阵、高德底图", "不直接导入未经核验的原始站点坐标"],
        ["智能能力", "可选 AI 策略建议、候选复核、结构化风险说明", "不让大模型直接生成并执行不可校验的路线"],
        ["交付方式", "Docker Compose、中文启动脚本、标准数据恢复、操作引导", "面向单县域业务域，支持后续扩展，不宣称多租户云 SaaS"],
    ], [2.4, 6.3, 5.95])
    add_heading(document, "1.3 文档导航", 2)
    add_table(document, "表 1-2 章节导航", ["章节", "阅读重点"], [
        ["第 2 章", "祁东县场景、参考资料、设施基线、产业货流与 AHP-GIS 依据"],
        ["第 3 章", "业务痛点、角色、需求、流程与验收口径"],
        ["第 4 章", "平台页面功能、操作入口与截图证据"],
        ["第 5 章", "数据域、空间校准、标准场景与路网服务边界"],
        ["第 6 章", "平台设计、实现细节、算法、AI、数据库、部署与风险"],
        ["第 7 章", "部署交付、测试结果、标准运行与验收清单"],
        ["附录", "数据字典、API 组、配置变量、参考资料与图表索引"],
    ], [2.4, 12.25])
    add_source_note(document, "阅读建议：先看第 2、3 章理解业务，再按第 4 章截图完成操作；需要排查部署或解释算法时阅读第 6、7 章和附录。")

    # Chapter 2 — reference-grounded Qidong baseline.
    add_heading(document, "2  祁东县场景与资料依据", 1)
    add_heading(document, "2.1 参考资料的使用方式", 2)
    add_body(document, "新增资料不是平台运行时的原始数据库，而是场景建模、设施规划、政策约束和数据校准的依据。平台将资料分成三层：研究材料用于解释业务背景和候选节点；汇总表格用于形成县域规模、乡镇需求与设施数量基线；平台业务数据用于保证操作过程可重复。原始材料中存在重复空行、版本差异和部分坐标质量问题，因此系统只吸收经过核对的汇总结论，未经核验的逐点坐标不直接进入运行场景。")
    add_table(document, "表 2-1 参考资料与系统用途", ["资料类型", "代表文件", "在平台中的用途"], [
        ["设施研究", "第四章（优化版）、第一版：祁东县城乡物流设施体系研究", "县—乡—村三级设施基线、覆盖率、设施缺口"],
        ["AHP-GIS 选址", "第三章基于 AHP-GIS 耦合模型的乡镇仓储分拨站选址适宜性前置筛选", "指标、权重、30 分钟服务圈候选排序"],
        ["产业与经济", "第一部分(1)", "产业货流、季节性、冷链和干线方向的业务解释"],
        ["行业摸排", "系统摸排祁东县物流行业整体发展水平", "道路、铁路、运量、冷链和行业约束"],
        ["统计表格", "人口密度、特色产业、乡镇需求与成本分析表", "乡镇分层、需求排序、成本与服务节点校准"],
        ["网点表格", "县级设施、乡镇站点、农村寄递服务站公示", "数量核对；逐点坐标仅在质量复核后使用"],
        ["政策资料", "政策梳理.docx", "跨部门协同、资金、土地、数据与人才风险"],
        ["路径规划参考", "实时路径规划优化系统代码/手册 PDF", "借鉴策略白名单、动态重分配、缓存和事件监控思想"],
    ], [2.6, 6.4, 5.65])
    add_heading(document, "2.2 县域基础画像", 2)
    add_body(document, "研究材料显示，祁东县具有县域核心仓配、工业园区货源和乡村农产品集散并存的复合物流结构。系统采用“县级枢纽—乡镇节点—村级服务站”的空间组织方式，将产业、人口、道路和设施数量作为调度解释层，而不是把规划研究中的总量直接当作当日订单。")
    add_table(document, "表 2-2 祁东县场景基线（研究资料核对值）", ["指标", "基线值", "对系统设计的影响"], [
        ["行政与人口", "24 个乡镇/街道、367 个村/社区；户籍人口约 103.6 万，常住人口约 72.62 万", "支持乡镇分层、服务半径和需求覆盖分析"],
        ["交通条件", "道路总里程 3,836.743 km；高速约 82.89 km；县域道路密度约 1.93 km/km²", "路网矩阵与县域边界需要分离展示"],
        ["县级设施", "5 处（4 处运营、1 处建设中）", "形成县域配送中心与转运枢纽层"],
        ["乡镇设施", "24 个乡镇物流站、24 个乡镇电商站；田头冷库约 95 座", "冷链兼容和多级集配是核心约束"],
        ["村级网络", "300 个村级寄递服务站，覆盖约 81.74%（300/367）", "识别约 20% 的村级覆盖缺口"],
        ["协同配送", "共配站约 197 个；集中冷库能力约 4 万吨", "支持共配、冷链和运力复用的业务解释"],
        ["行业运量", "2025 年道路货运量约 409.06 万吨，周转量约 40,885.43 万吨公里", "强调干线—县内支线联动，不把样本订单当作全县运量"],
    ], [3.0, 5.4, 6.25])
    add_heading(document, "2.3 产业货流与业务约束", 2)
    add_table(document, "表 2-3 重点产业货流对调度规则的影响", ["货流来源", "研究资料特征", "平台约束映射"], [
        ["黄花菜与油茶", "黄花菜产量和加工季节集中；油茶采收期短，预冷不足会增加损耗", "冷链车辆优先、季节性风险提示、时间窗约束"],
        ["果蔬", "蔬菜、柑橘和脆枣存在短保鲜期与较高损耗风险", "服务等级、冷链兼容、预计延误和风险说明"],
        ["生猪与农资", "生猪运输存在损耗；农资和饲料具有批量、时限和返程特征", "重量、时窗、车辆班次与双向配送解释"],
        ["建材", "月度峰值高、外运占比高，回程空驶率约 72%", "成本、里程、空驶率和车辆利用率对比"],
        ["纺织与机电", "小批量、多批次、园区集中，部分货流为零担与专用通道", "节点聚合、车辆类型、线路复用和方案评分"],
    ], [3.0, 6.25, 5.4])
    add_heading(document, "2.4 需求、设施与坐标数据口径", 2)
    add_body(document, "乡镇需求与成本分析表汇总需求约 5,557.1 吨（不同版本表格因舍入存在约 0.6 吨差异），它是规划研究期间的聚合量，不是平台某一天的订单量。标准场景将该规模按 DEMO_LOAD_RATIO 归一化为约 8.3 吨的 16 条标准订单，便于在常规运行环境中完成完整求解。农村寄递站点表的部分坐标呈单调序列，存在录入或导出质量疑点；平台因此使用空间校准后的节点，并在资料说明中保留风险，而不是用未经核验的点位制造“精确地图”。")
    add_callout(document, "版本口径", "AHP-GIS 资料有 60 分钟服务圈的第一版和 30 分钟服务圈的优化版。平台运行时引用 30 分钟版本：石亭子、鸟江、砖塘、河洲和过水坪位于前列；60 分钟版本仅作为历史研究对照，不能与运行时排序混用。", fill="#FFF5DD", accent="#B97915")
    add_table(document, "表 2-4 平台运行时采用的 AHP-GIS 权重", ["指标", "含义", "权重"], [
        ["C1 坡度", "建设与车辆通行适宜性", "0.109"],
        ["C2 高程", "地形与基础设施建设条件", "0.055"],
        ["C3 道路密度", "可达性与干线衔接", "0.359"],
        ["C4 服务半径", "30 分钟服务圈覆盖能力", "0.180"],
        ["C5 农业覆盖", "特色产业与农产品集散支撑", "0.223"],
        ["C6 工业布局", "工业园区及产业货源支撑", "0.074"],
    ], [3.2, 8.9, 2.0])
    add_table(document, "表 2-5 30 分钟服务圈候选站点排序（运行时参考）", ["序位", "候选乡镇", "得分", "备注"], [
        ["1", "石亭子镇", "7.631", "综合适宜性最高"],
        ["2", "鸟江镇", "7.544", "道路与产业综合较强"],
        ["3", "砖塘镇", "7.541", "农产品集散优势"],
        ["4", "河洲镇", "7.499", "服务半径表现较好"],
        ["5", "过水坪镇", "7.484", "油茶等产业支撑"],
        ["6", "金桥镇", "7.480", "产业与道路条件均衡"],
        ["7", "白地市镇", "7.406", "节点覆盖能力较好"],
        ["8", "双桥镇", "7.321", "候选备选"],
        ["9", "归阳镇", "7.209", "干线衔接与产业并重"],
        ["15", "凤歧坪镇", "3.323", "偏远保障候选，不能简单淘汰"],
    ], [1.4, 4.4, 2.0, 6.85])

    # Chapter 3 — requirements and operating model.
    add_heading(document, "3  业务需求与产品方案", 1)
    add_heading(document, "3.1 需要解决的业务问题", 2)
    add_numbered(document, [
        "订单、车辆和网点分散在表格或个人经验中，难以快速判断当天运力是否足够。",
        "县乡村节点多、道路条件复杂，直线距离不能代表真实行驶距离，路线结果需要道路矩阵和坐标口径支撑。",
        "冷链车辆数量有限，故障或紧急订单发生后，人工重新分配容易破坏时间窗和容量约束。",
        "方案生成后缺少审批、下发、驾驶员接单与完成归档，调度结果难以形成可执行任务。",
        "大模型、地图与路网服务存在鉴权和额度风险，系统必须能解释是否使用 AI、是否使用真实路网以及何时降级。",
    ])
    add_heading(document, "3.2 目标与成功标准", 2)
    add_table(document, "表 3-1 目标—功能—证据映射", ["目标", "功能实现", "验收证据"], [
        ["可计算", "订单/车辆/节点结构化；OR-Tools 约束求解", "调度页显示方案、约束与运行编号"],
        ["可解释", "透明评分、策略白名单、AI 复核边界", "方案卡片、评分分项和 AI 状态说明"],
        ["可执行", "审批、下发、司机任务、状态机", "执行中心与驾驶员任务页面"],
        ["可恢复", "标准场景、一键恢复业务数据", "恢复后订单、车辆、事件回到基线"],
        ["可移交", "Docker、一键脚本、配置保存、操作引导", "部署人员按说明启动脚本和浏览器"],
        ["可审计", "用户、配置、求解、事件与执行日志", "审计页与历史方案页"],
    ], [2.7, 6.0, 5.95])
    add_heading(document, "3.3 角色与权限模型", 2)
    add_table(document, "表 3-2 用户角色与权限边界", ["角色", "可见模块", "关键权限"], [
        ["系统管理员", "全部模块与设置", "账号、AI/地图配置、数据恢复、审计、审批与下发"],
        ["运营用户/调度员", "运营、调度、事件、执行、数据、历史", "创建方案、处理事件、提交审批；不能读密钥和系统审计"],
        ["驾驶员", "运营总览、本人执行任务、动态消息", "接单、发车、完成本人任务；不能修改方案和配置"],
        ["查看者（预留）", "只读运营与报告", "查看指标与报告，不改变业务状态"],
    ], [3.1, 5.5, 6.05])
    add_heading(document, "3.4 业务闭环与操作顺序", 2)
    add_body(document, "日常操作遵循“登录—查看总览—校验基础数据—生成调度—比较并提交—审批下发—驾驶员执行—事件重调度—归档复盘”的顺序。系统内置操作引导只负责定位入口和解释下一步，不会替用户自动审批或下发；标准数据恢复会保留账号、AI 配置和高德配置，避免误操作后无法回到稳定状态。")
    add_figure(document, "20-platform-logic-image2.png", "图 3-1 祁智运平台逻辑结构设计", 14.0)
    add_figure(document, "21-function-framework-image2.png", "图 3-2 祁智运平台功能框架图", 14.0)
    add_figure(document, "10-business-loop.png", "图 3-3 祁智运平台业务决策闭环", 14.0)

    # Chapter 4 — page-by-page product evidence.
    add_heading(document, "4  平台功能与界面设计", 1)
    add_body(document, "本章按实际菜单和操作顺序说明平台功能，并以运行截图作为实现证据。截图中的订单、车辆和节点采用统一的标准场景；外部服务状态、温度和路网耗时会随网络与配额变化。")
    add_heading(document, "4.1 注册、登录与会话安全", 2)
    add_body(document, "登录页是所有角色的统一入口。公开注册只创建最低权限运营用户，密码必须满足长度、大小写、数字和特殊字符组合要求；登录失败会计数并触发临时锁定，成功登录后由后端签发带角色信息和过期时间的 JWT。前端路由隐藏不是安全边界，所有业务 API 仍由后端 RBAC 再校验。")
    add_figure(document, "01-login.png", "图 4-1 统一登录入口", 14.0)
    add_figure(document, "02-register.png", "图 4-2 运营用户注册页面", 14.0)
    add_figure(document, "22-login-flow-image2.png", "图 4-3 登录与注册流程图", 14.0)
    add_heading(document, "4.2 运营总览与需求预测", 2)
    add_body(document, "运营总览把今日预测需求、建议运力、路网矩阵来源和预计准时率集中到首屏，并以未来 7 日预测、乡镇需求排行、冷链需求、风险预警和县域网络地图辅助判断。样本不足时，系统明确标识星期季节基线，不把统计基线包装成真实机器学习精度。管理员和驾驶员看到同一业务入口，但指标按角色和可见数据范围呈现，驾驶员只看到与本人任务有关的执行信息。")
    add_figure(document, "03-overview.png", "图 4-4 县域物流智能决策驾驶舱（管理员视角）", 14.0)
    add_heading(document, "4.3 场景、订单与车辆基础数据", 2)
    add_body(document, "基础数据中心提供订单台账、车辆档案、节点与空间校准入口。标准场景 qidong-real-001 包含 16 条标准订单、5 辆车辆和 18 个空间节点；其中冷链订单与冷链车辆的匹配由后端约束校验，不依赖前端勾选。场景页用于确认数据版本，历史页用于查阅已完成的求解运行。")
    add_figure(document, "08-master-data.png", "图 4-5 基础数据中心订单台账", 14.0)
    add_figure(document, "18-scenarios.png", "图 4-6 稳定业务场景与版本信息", 14.0)
    add_heading(document, "4.4 智能调度与路径规划", 2)
    add_body(document, "调度中心先读取场景与空间快照，再请求 OSRM 或降级矩阵，最后由 OR-Tools 生成可行路线。界面同时展示车辆路线、节点顺序、装载率、里程、服务率和优化前后对比；AI 仅在启用且调用成功时参与策略建议或候选复核。建议方案必须通过容量、时间窗、冷链、班次和可用性校验，才能提交审批。")
    add_figure(document, "04-dispatch.png", "图 4-7 智能调度中心与路线方案", 14.0)
    add_heading(document, "4.5 动态事件与实时重调度", 2)
    add_body(document, "动态事件页内置固定车辆故障事件，便于重复验证“识别—约束重建—增量求解—方案下发”的路径。处置前冻结已执行任务；处置时移除故障车辆并重构剩余任务；处置完成后，故障标记、左侧消息数和右侧结构化说明同步更新。该设计避免把一段长文本当作处置结果，也让每次事件有明确的前后状态。")
    add_figure(document, "05-events.png", "图 4-8 动态事件与实时重调度页面", 14.0)
    add_heading(document, "4.6 审批、下发与驾驶员执行", 2)
    add_body(document, "调度结果只有在提交并通过审批后，才会拆分为驾驶员任务。执行中心展示方案状态、车辆任务、停靠顺序和完成进度；驾驶员账号只允许查看和更新本人任务的合法状态，不能修改系统配置或他人的任务。每一次状态变更记录操作人、时间、前后状态和备注，完成后进入历史与报告。")
    add_figure(document, "06-execution.png", "图 4-9 调度审批与执行中心", 14.0)
    add_heading(document, "4.7 算法评估、历史与准备度", 2)
    add_body(document, "算法评估以标准场景和相同路网口径比较业务基线与优化方案，保存里程、成本、车辆数、服务率、未分配订单、耗时和指纹。历史页用于回看运行版本，准备度页用于检查数据、算法、执行环节、系统保障和交付材料是否齐备。该页的分数是可解释的准备度指标，不替代真实运营 KPI。")
    add_figure(document, "07-experiments.png", "图 4-10 算法对照实验室", 14.0)
    add_figure(document, "19-history.png", "图 4-11 方案运行历史", 14.0)
    add_figure(document, "17-readiness.png", "图 4-12 系统准备度检查", 14.0)
    add_heading(document, "4.8 AI、地图、用户与审计设置", 2)
    add_body(document, "服务配置采用连接列表。管理员可新增、编辑、测试、启用或停用 AI 连接，API Key 只在数据库中以 Fernet 密文保存，界面仅返回首尾掩码；高德 Web Key 与安全密钥采用同样的加密和掩码原则。系统设置页还包含用户管理、审计日志和地图状态，方便定位 401、域名白名单、额度或代理配置问题。标准数据恢复只重置业务样本，不覆盖这些服务连接与账号设置。")
    add_figure(document, "09-settings-ai.png", "图 4-13 AI 服务连接管理", 14.0)
    add_figure(document, "14-settings-map.png", "图 4-14 高德地图服务配置", 14.0)
    add_figure(document, "15-settings-users.png", "图 4-15 用户与角色管理", 14.0)
    add_figure(document, "16-settings-audit.png", "图 4-16 审计日志", 14.0)
    add_heading(document, "4.9 操作引导与标准恢复", 2)
    add_body(document, "操作引导以五步路径提示使用者下一步点击哪里：运营总览、生成调度、查看方案、处理事件、完成执行。引导不会自动触发求解和下发，顶部入口可随时重新打开。若订单状态被改乱、故障事件已经处理或业务数据需要恢复，使用“一键恢复标准场景数据”即可回到与本文件截图一致的业务基线；恢复前会先生成备份，不会清理 Docker 数据卷。")

    # Chapter 5 — data and spatial governance.
    add_heading(document, "5  数据治理与空间一致性", 1)
    add_heading(document, "5.1 数据域与来源层次", 2)
    add_table(document, "表 5-1 数据域、主要字段与来源", ["数据域", "主要对象", "来源/治理动作"], [
        ["身份与审计", "用户、角色、会话、操作日志", "数据库事务；密码哈希；配置和权限动作留痕"],
        ["基础业务", "订单、车辆、节点、乡镇映射", "稳定场景；导入前校验重量、时窗、引用和冷链"],
        ["空间治理", "GCJ-02 坐标、来源、精度、校准记录", "人工核验点位；修订表；生成空间快照"],
        ["路网与求解", "距离矩阵、时间矩阵、策略、约束和输出", "OSRM 或明确标识的估算矩阵；缓存与指纹"],
        ["事件与执行", "故障、紧急订单、方案状态、司机任务", "状态机；冻结已执行任务；事件与执行审计"],
        ["外部服务", "AI/高德连接、测试状态、密文 Key", "Fernet 加密入库；接口掩码；单一启用连接"],
    ], [2.6, 5.5, 6.0])
    add_heading(document, "5.2 研究数据与运行数据的区分", 2)
    add_body(document, "研究资料回答“祁东县物流网络应该如何规划”，运行数据回答“本次调度要给哪些订单和车辆安排路线”。例如乡镇需求分析表的约 5,557.1 吨是聚合研究量，16 条标准订单的约 8.3 吨是可重复运行样本；两者都保留，但用途不同。文档、接口和页面应始终显示数据版本或来源，以免把规划量误读为当天承运承诺。")
    add_heading(document, "5.3 坐标、底图与路网矩阵", 2)
    add_table(document, "表 5-2 空间服务分工", ["组件", "职责", "故障时的表现"], [
        ["高德 JSAPI", "浏览器底图、行政地理要素和可视化", "底图不可用时显示离线边界或提示配置问题"],
        ["空间校准", "维护节点坐标来源、精度、修订与快照", "点位不完整时阻止静默写入，提示质量问题"],
        ["OSRM Table", "服务端距离/时间矩阵，支持真实道路路线", "超时/不可达时使用估算矩阵并记录降级原因"],
        ["OR-Tools", "在矩阵和业务约束上求解车辆路径", "无可行解时返回约束冲突，不生成假路线"],
        ["离线 GeoJSON", "弱网场景提供县域边界和基本地图参照", "不代替真实道路矩阵，不宣称导航级精度"],
    ], [2.8, 6.3, 5.0])
    add_callout(document, "坐标质量", "农村寄递服务站表中部分点位呈单调序列，本平台仅用于数量核对，不直接作为地图标点。运行节点由空间校准和标准场景显式维护，并在求解结果中保存空间快照编号。", fill="#FFF5DD", accent="#B97915")
    add_heading(document, "5.4 标准场景数据映射", 2)
    add_body(document, "为了避免每次打开页面结果漂移，系统把 24 个乡镇聚合需求映射到 16 个运行节点，并固定订单权重、冷链属性、时间窗、车辆容量和单位成本。该映射是“研究数据到可计算样本”的透明转换，不是对每个村级站点的精确还原；正式生产接入时，应将真实订单和经审核的站点坐标另建场景版本。")
    add_table(document, "表 5-3 数据质量检查清单", ["检查项", "通过条件", "不通过时的动作"], [
        ["订单字段", "重量>0、节点存在、时间窗可解析", "导入预检失败，不写入业务表"],
        ["车辆字段", "容量>0、状态可用、冷链能力明确", "标记不可用或要求修正"],
        ["坐标", "经纬度在祁东边界内，来源和精度可追溯", "进入校准队列，不参与求解"],
        ["路网矩阵", "矩阵维度与节点一致，距离/时间非负", "回退估算并记录 provider/fallback"],
        ["结果版本", "场景、空间快照、路网元数据和指纹齐全", "禁止提交审批"],
    ], [2.7, 6.1, 5.3])

    add_heading(document, "6  祁智运县域物流智能决策平台设计与实现", 1)
    add_body(document, "祁智运面向祁东县黄花菜等特色农产品的城乡双向配送场景，将订单、车辆、服务网点、道路矩阵和执行状态纳入同一业务空间。平台并非单纯展示地图或调用大模型，而是以“可信数据—约束求解—业务审批—任务执行—异常反馈”为主线，把县域物流中分散的经验调度转化为可计算、可复核、可追溯的数字化流程。")
    add_body(document, "平台采用确定性优化与生成式人工智能协同的混合架构。OR-Tools 负责容量、时间窗、冷链兼容和车辆班次等硬约束；OSRM 提供道路距离与行驶时间矩阵，高德地图负责前端真实底图；大模型在启用且连接正常时参与业务目标理解、策略建议和安全候选复核，但不得修改硬约束，也不得选择不可行路线。即使外部 AI 或地图服务不可用，平台仍可依靠确定性求解与离线县域边界完成核心业务。")

    add_heading(document, "6.1 平台目标与设计原则", 2)
    add_heading(document, "6.1.1 建设目标", 3)
    add_body(document, "平台建设目标是形成一套可在常规运行环境中稳定运行的县域物流智能决策系统。它统一维护订单、车辆、网点和坐标数据，依据真实路网或明确标识的降级矩阵生成配送方案，并把调度结果继续推进到审批、下发和驾驶员执行阶段。管理人员可从运营总览判断当日需求与运力，调度人员可比较路线与成本，驾驶员可处理个人任务，系统管理员则负责数据、账号、AI 服务和地图服务治理。")
    add_body(document, "围绕“最后一公里成本高、冷链资源有限、乡村网点分散、突发事件响应慢”等问题，系统给出可量化的里程、成本、服务率、车辆数、装载率和预计延误指标。每次求解记录场景版本、道路来源、候选策略、评分分项和运行编号，使结果可以重现，而不是只能看到一张无法解释的线路图。")

    add_heading(document, "6.1.2 设计原则", 3)
    add_numbered(document, [
        "业务约束优先。容量、时间窗、冷链、车辆班次和车辆可用状态由确定性求解器统一校验，任何智能建议都不能绕过硬约束。",
        "数据与地图同源。网点坐标采用 GCJ-02 口径并保留校准记录，路线展示与道路矩阵分别标识来源，避免“地图看似正确而求解仍用直线距离”。",
        "智能能力可降级。AI、OSRM 或高德服务不可用时，系统展示明确的降级状态并继续完成可行求解，不用虚构结果掩盖外部依赖故障。",
        "全过程可追溯。注册、登录、配置、数据发布、方案审批、任务下发与执行动作均进入审计链，关键结果使用场景指纹和快照固化。",
        "便捷交付。采用 Docker Compose 和中文一键脚本完成启动、停止、状态检查与标准场景恢复，部署人员无需单独安装数据库或运行时。",
        "最小权限与密钥安全。公开注册只生成普通运营账号，管理员权限不可自授；API Key 和地图安全密钥加密入库，前端和接口只显示掩码。",
    ])

    add_heading(document, "6.2 业务架构与运行逻辑", 2)
    add_heading(document, "6.2.1 县域物流业务闭环", 3)
    add_body(document, "系统以一个激活的业务场景作为当日计算基线。订单台账提供重量、服务网点、时间窗、冷链与优先级；车辆档案提供载重、车型、冷链能力、班次与成本参数；网点档案提供县域配送中心、乡镇和村级服务点的坐标。数据质量检查通过后，平台读取道路矩阵，结合需求预测和风险识别形成求解实例。")
    add_body(document, "求解阶段可生成多种安全候选方案，对各候选计算服务率、里程改善、成本改善、车辆利用和空驶控制分数。启用 AI 时，大模型只能在安全候选集合内给出复核意见；未启用或调用失败时，系统按透明评分直接选取方案。方案生成后仍需经过提交、审批、下发，驾驶员再依次确认接单、发车和送达。车辆故障等事件发生时，系统冻结已执行任务，只对剩余任务重构求解实例。")

    add_heading(document, "6.2.2 用户角色与权限边界", 3)
    add_table(
        document,
        "表 6-1 平台角色与主要权限",
        ["角色", "核心职责", "主要权限", "限制"],
        [
            ["系统管理员", "系统治理", "用户、配置、场景、审计与全部业务", "敏感密钥仍不可查看明文"],
            ["调度员", "方案生成与下发", "调度、事件、方案、执行与运营数据", "不能管理系统密钥和管理员账号"],
            ["驾驶员", "运输执行", "查看个人任务并确认接单、发车、完成", "不能查看他人任务和系统配置"],
            ["查看者", "运营观察", "只读查看被授权的数据与结果", "不能生成、审批或修改业务数据"],
        ],
        [2.3, 3.2, 5.1, 4.0],
    )

    add_heading(document, "6.3 平台功能设计", 2)
    add_heading(document, "6.3.1 注册、登录与账号安全", 3)
    add_body(document, "登录页面采用独立的企业身份入口，支持会话存储与“在此设备保持登录”。连续登录失败达到阈值后，账号会被临时锁定；成功登录将清空失败次数并记录最近登录时间和 IP。JWT 令牌携带用户与角色信息，前端路由和后端接口同时校验权限，避免仅隐藏按钮却未限制接口。")
    add_body(document, "公开注册页面要求登录账号、姓名或称呼、强密码、二次确认和账号数据使用规范确认。账号必须以字母开头，仅允许字母、数字、点、短横线和下划线；密码至少 12 位，并同时包含大小写字母、数字和特殊字符。后端固定将公开注册账号设置为最低权限运营用户，任何额外提交的 role 或管理员字段都会被拒绝。注册成功后自动签发会话，重复账号返回冲突提示，注册动作同时写入审计日志。")

    add_heading(document, "6.3.2 运营总览与需求预测", 3)
    add_body(document, "运营总览把当日需求、建议运力、道路矩阵来源和预计准时率放在首屏，并组合未来 7 日需求预测、乡镇需求排行、冷链需求、风险预警与县域配送网络。历史样本不足 28 天时，系统使用带星期季节因子的统计基线，并明确展示样本量和降级状态；达到门槛后才允许训练按时间切分验证的机器学习模型，禁止用未来数据泄漏提高表面精度。")
    add_body(document, "地图同时展示祁东县行政边界、配送中心、乡镇村级节点和线路。高德 JavaScript Key 只承担浏览器底图渲染，OSRM 或后续独立 Web 服务 Key 承担服务端道路矩阵，两者权限和用途严格分离。右侧“生成今日调度”进入调度中心，既降低操作层级，也与内置操作引导的第一步一致。")

    add_heading(document, "6.3.3 基础数据与空间校准", 3)
    add_body(document, "基础数据中心统一管理订单台账、车辆档案、网点站点、空间校准和数据质量。订单支持搜索、状态与优先级筛选、冷链筛选、新建、编辑、删除和批量导入；导入文件先进入预检批次，字段、引用、时间窗、重量和冷链运力全部校验后才能提交，校验失败不会写入半条业务数据。")
    add_body(document, "空间校准记录 GCJ-02 坐标、来源、状态、精度、校准人员与时间，并把每次坐标变更写入修订表。系统为求解生成空间坐标快照，后续路线、实验和报告引用同一快照，解决坐标修改后历史结果无法复现的问题。")

    add_heading(document, "6.3.4 智能调度与路线可视化", 3)
    add_body(document, "智能调度中心是平台的核心工作区。顶部地图按车辆区分路线颜色，显示配送顺序、冷链属性、节点、道路来源和坐标状态；右侧路线卡片展示车辆、节点数、装载率、里程和停靠顺序。下方控制台展示订单量、货量、冷链订单、准时率、综合评分以及优化前后里程、空驶率、成本和车辆数对比。")
    add_body(document, "调度求解默认读取标准业务场景和 OSRM 道路矩阵，并对容量、时间窗、车辆班次、冷链兼容与车辆可用性进行硬约束校验。系统不会因为大模型给出一段解释就宣称已完成 AI 求解；界面分别显示确定性求解结果、AI 是否参与策略建议或复核、道路矩阵是否降级以及动态评分证据。")

    add_heading(document, "6.3.5 动态事件与实时重调度", 3)
    add_body(document, "动态事件页面内置一套固定车辆故障事件：冷链车辆在既定网点发生故障，受影响订单被重新识别，同时加入一条紧急冷链订单。故障图标、事件告警、受影响订单数、最大延误和冷链约束显示在同一页面，便于操作人员先确认事件，再启动处置。")
    add_body(document, "处置链包括事件识别、约束重建、增量求解和方案下发。系统保留已经完成或正在执行的任务快照，将故障车辆从可用集合中移除，重新分配剩余订单；完成处置后，故障图标与未处理消息同步清除，页面以结构化条目展示策略、约束、路网、快照和风险，不再堆叠混乱的长文本。")

    add_heading(document, "6.3.6 审批、下发与驾驶员执行", 3)
    add_body(document, "求解结果并不直接变成车辆任务。管理员或调度员先把可行 run_id 提交为执行方案，管理员审核后再下发；系统按每条路线生成独立驾驶员任务，固化路线、停靠顺序和指标快照。驾驶员登录后只能读取分配给自己的任务，并依次执行接单、发车和完成。非法状态跳转返回冲突，越权访问返回禁止。")
    add_body(document, "执行中心以六个状态节点展示“提交审批—审核通过—任务下发—司机接单—配送执行—完成归档”，每次动作记录人员、时间、前后状态和备注。所有车辆任务完成后，执行方案自动汇总为完成状态，并可下载自包含的执行报告。")

    add_heading(document, "6.3.7 算法评估与方案追溯", 3)
    add_body(document, "算法评估页面在同一场景、同一道路矩阵和同一统计口径下，对业务基线与 OR-Tools 推荐方案进行比较。一次实验固定运行标准负荷、旺季峰值和冷链压力三个案例，保存里程、成本、车辆数、服务率、未分配订单、求解耗时、求解器、场景指纹和路网来源。")
    add_body(document, "业务基线用于量化改进空间，不被描述为数学最优解。实验历史和方案历史分别保留运行编号、版本和时间，用户可以下载自包含报告，避免只保留一张无法复核的截图。标准基线在三个压力案例中均能保持 100% 服务率，并在标准负荷下显著降低里程和成本。")

    add_heading(document, "6.3.8 AI 服务与地图服务配置", 3)
    add_body(document, "系统设置采用连接列表而非单一裸表单。管理员可新增、编辑、测试、启用或停用多条 AI 服务连接，任意时刻只允许一个默认启用连接。保存 API Key 后数据库仅存储 Fernet 密文，接口只返回首尾掩码；编辑模型、地址或温度时可留空保留原密钥。连接测试记录最近状态、时间和耗时，鉴权失败会原样提示服务端状态但不会泄露完整密钥。")
    add_body(document, "地图服务同样采用连接治理模式，Web Key 与安全密钥分别加密保存。浏览器运行时只获得 Web Key、地图样式和同源安全代理地址，安全密钥由服务端代理追加，不返回前端。默认连接随一键恢复操作保留，防止恢复业务基线时误删使用者已经配置的外部服务。")

    add_heading(document, "6.3.9 操作引导与标准状态恢复", 3)
    add_body(document, "平台内置五步操作引导，依次定位运营总览、生成调度、调度方案、异常处置和执行闭环。引导只切换页面与高亮控件，不自动触发求解、审批或下发，避免误操作。首次登录自动出现，完成后写入浏览器本地状态，顶部入口可随时重新启动。")
    add_body(document, "标准状态恢复把订单、车辆、网点、算法实验和固定故障事件恢复到可重复的业务基线，同时明确保留 AI Key、高德地图连接和账号数据。恢复脚本先生成数据库备份，再调用受控恢复流程；因此即使使用者改变了订单状态或处置了故障，也可回到文档与截图一致的标准状态。")

    add_heading(document, "6.4 平台技术架构设计", 2)
    add_heading(document, "6.4.1 分层架构", 3)
    add_body(document, "平台采用前后端分离和容器化架构。接入层由 Nginx 提供静态资源和 /api 反向代理；应用层使用 Vue 3、TypeScript、Vue Router、ECharts 和高德 JSAPI；服务层使用 FastAPI 实现 REST 接口，Celery Worker 处理异步优化任务；智能层组合 OR-Tools、OSRM 路网矩阵和可选大模型；数据层使用 PostgreSQL 16、PostGIS 3.4 和 Redis。")
    add_body(document, "该架构没有为追求概念复杂度而拆分大量微服务，而是在单一业务边界内保持模块化。同步请求适合登录、查询、配置和小规模求解；耗时任务通过 Redis 投递给 Worker，前端使用任务 ID 查询状态。数据库迁移由独立 migrate 容器在业务服务启动前执行，避免应用实例并发迁移。")
    add_figure(document, "11-technical-architecture.png", "图 6-12 祁智运平台分层技术架构", 14.5)
    add_figure(document, "23-microservice-architecture-image2.png", "图 6-0 祁智运平台微服务架构图", 14.5)

    add_heading(document, "6.4.2 核心 API 设计", 3)
    add_body(document, "API 统一使用 /api/v1 前缀，输入输出采用 JSON，文件导入使用 multipart/form-data。认证接口公开，业务接口要求 Bearer JWT；管理员、调度员、驾驶员和查看者的角色依赖后端权限函数校验。错误响应区分未认证、禁止、冲突、参数错误与外部服务故障，便于前端给出清晰反馈。")
    add_table(
        document,
        "表 6-2 核心 API 分组",
        ["接口分组", "代表路径", "主要用途"],
        [
            ["认证", "POST /auth/register；POST /auth/login", "注册、登录、当前用户与退出"],
            ["调度", "POST /optimize；GET /runs", "生成候选方案并查询历史"],
            ["事件", "POST /events/reoptimize", "车辆故障与紧急订单重调度"],
            ["执行", "POST /dispatch/plans；GET /driver/assignments", "审批、下发与驾驶员任务"],
            ["基础数据", "GET/POST /master-data/*", "订单、车辆、网点、导入与质量治理"],
            ["外部服务", "GET/POST /ai/configs；/map/configs", "连接管理、测试和启停"],
        ],
        [2.3, 5.8, 6.55],
    )

    add_heading(document, "6.4.3 安全与可审计设计", 3)
    add_body(document, "用户密码使用 PBKDF2-HMAC-SHA256 加随机盐并迭代 240000 次保存，不存储可逆密码。注册与管理员建号共用强密码策略；登录失败计数和临时锁定减少撞库风险。JWT 使用服务端密钥签名并设置有效期，生产环境要求不少于 32 字节的随机密钥。")
    add_body(document, "AI Key、地图 Web Key 和安全密钥由独立 Fernet 密钥加密。密钥轮换需要先解密旧密文、换钥重加密并验证，不能在已有配置后直接替换环境变量。审计日志记录类别、动作、结果、用户、IP、User-Agent、资源和结构化明细；普通用户无法读取系统审计和配置。")
    add_table(
        document,
        "表 6-3 关键安全控制",
        ["控制对象", "实现机制", "安全结果"],
        [
            ["密码", "PBKDF2 + 随机盐 + 强度策略", "数据库泄露时不能直接还原明文"],
            ["会话", "JWT 有效期 + 前后端双重鉴权", "路由隐藏与接口权限保持一致"],
            ["公开注册", "固定最低权限 + extra=forbid", "不能通过附加字段自授管理员"],
            ["外部密钥", "Fernet 加密 + 掩码输出", "浏览器与接口不返回完整密钥"],
            ["关键操作", "结构化审计日志", "账号、配置、数据和执行动作可追溯"],
        ],
        [2.8, 6.0, 5.85],
    )

    add_heading(document, "6.5 优化算法模型与 AI 协同", 2)
    add_heading(document, "6.5.1 问题定义", 3)
    add_body(document, "平台求解的是带容量、时间窗、车辆班次、冷链兼容和车辆可用性的多车辆路径问题。设配送中心为节点 0，订单节点集合为 N，车辆集合为 K。每个订单 i 具有需求重量 qᵢ、服务时间 sᵢ、时间窗 [aᵢ,bᵢ] 和冷链属性 cᵢ；每辆车 k 具有容量 Qₖ、班次 [Aₖ,Bₖ]、单位里程成本 Cₖ 和冷链能力 rₖ。道路矩阵给出节点 i 到 j 的距离 dᵢⱼ 与时间 tᵢⱼ。")
    add_table(
        document,
        "表 6-4 主要符号定义",
        ["符号", "定义", "单位"],
        [
            ["N，K", "订单节点集合、可用车辆集合", "—"],
            ["qᵢ，Qₖ", "订单重量、车辆容量", "kg"],
            ["dᵢⱼ，tᵢⱼ", "道路距离、道路行驶时间", "km，min"],
            ["[aᵢ,bᵢ]", "订单服务时间窗", "min"],
            ["[Aₖ,Bₖ]", "车辆工作班次", "min"],
            ["cᵢ，rₖ", "订单冷链要求、车辆冷链能力", "0/1"],
            ["xᵢⱼₖ", "车辆 k 是否从 i 行驶到 j", "0/1"],
        ],
        [2.4, 9.0, 3.25],
    )

    add_heading(document, "6.5.2 目标函数与约束", 3)
    add_body(document, "基础目标是在全部硬约束可行的前提下，最小化行驶距离、车辆成本、未服务惩罚和超时惩罚的加权和。平台不使用大模型直接生成节点序列，而是把允许的策略参数限制在距离优先、成本优先与均衡策略白名单内，再由 OR-Tools Routing 求解。")
    add_equation(document, "min Z = ΣₖΣᵢΣⱼ dᵢⱼ·xᵢⱼₖ + λ₁·Cost + λ₂·Unserved + λ₃·Late", "6-1")
    add_equation(document, "Σᵢ qᵢ·yᵢₖ ≤ Qₖ，∀ k ∈ K", "6-2")
    add_equation(document, "Tⱼₖ ≥ Tᵢₖ + sᵢ + tᵢⱼ − M(1−xᵢⱼₖ)", "6-3")
    add_equation(document, "aᵢ ≤ Tᵢₖ ≤ bᵢ，Aₖ ≤ Tᵢₖ ≤ Bₖ", "6-4")
    add_equation(document, "cᵢ·yᵢₖ ≤ rₖ，∀ i ∈ N，k ∈ K", "6-5")
    add_body(document, "式（6-2）保证车辆装载量不超过容量；式（6-3）维持路径时间递推；式（6-4）同时约束订单时间窗和车辆班次；式（6-5）保证冷链订单只能分配给冷链车辆。车辆故障发生后，故障车辆从 K 中剔除，已执行节点固定不动，剩余节点重新求解。")

    add_heading(document, "6.5.3 道路矩阵与坐标一致性", 3)
    add_body(document, "OSRM Table 接口为全部节点生成距离和时间矩阵，相同坐标快照可命中缓存。若 OSRM 不可达、返回异常或坐标不完整，系统切换县乡道路估算矩阵，并把 requested_provider、source、fallback_used、fallback_reason、data_version 和 coordinate_count 写入结果。地图底图是否加载成功不能代替道路矩阵状态，二者在界面中独立展示。")

    add_heading(document, "6.5.4 AI 协同与透明评分", 3)
    add_body(document, "大模型参与两个可选阶段：一是把业务目标转换为白名单策略建议，二是在确定性求解器产生的可行候选集合内进行复核。AI 输出必须解析为结构化 JSON，非法字段直接忽略；若鉴权失败、超时、未启用或推荐结果与透明评分差距过大，系统拒绝 AI 选择并回退确定性最优候选。")
    add_figure(document, "12-ai-boundary.png", "图 6-13 大模型与确定性求解器的协同边界", 14.5)
    add_table(
        document,
        "表 6-5 调度方案透明评分结构",
        ["评分项", "权重", "计算依据"],
        [
            ["订单服务与约束满足", "40", "服务率、未分配订单和可行性"],
            ["里程改善", "25", "与同一道路矩阵业务基线的改善比例"],
            ["成本改善", "20", "与同一道路矩阵业务基线的改善比例"],
            ["车辆利用", "10", "使用车辆数与可用车辆数的合理程度"],
            ["空驶控制", "5", "空驶率阈值分段"],
        ],
        [4.2, 2.0, 8.45],
    )

    add_heading(document, "6.6 数据库架构设计", 2)
    add_heading(document, "6.6.1 数据域划分", 3)
    add_body(document, "PostgreSQL 将平台数据划分为身份与审计、基础数据、场景治理、优化运行、动态事件、执行闭环和外部服务配置七个数据域。业务表通过主外键和检查约束保持一致性，运行结果中的复杂路线、评分证据和路网元数据使用 JSONB 保存，以兼顾结构化查询和求解输出扩展。PostGIS 为后续空间查询和路网本地化预留能力。")
    add_table(
        document,
        "表 6-6 核心数据实体",
        ["数据域", "核心表", "主要内容"],
        [
            ["身份与审计", "users，audit_logs", "角色、锁定、登录信息与操作留痕"],
            ["基础数据", "orders，vehicles，logistics_nodes", "订单、车辆、网点和业务状态"],
            ["空间治理", "node_coordinate_revisions", "坐标修改前后值、来源、精度与人员"],
            ["场景治理", "scenario_versions，scenario_transitions", "不可变版本、指纹、审批和激活"],
            ["优化与事件", "optimization_runs，dispatch_events", "路线、指标、候选证据和故障事件"],
            ["执行闭环", "dispatch_plans，driver_assignments", "审批状态、路线快照与驾驶员进度"],
            ["外部服务", "ai_provider_configs，map_provider_configs", "加密凭据、启停和连接测试"],
        ],
        [2.6, 5.1, 7.0],
    )
    add_figure(document, "24-system-er-image2.png", "图 6-0 祁智运系统总体 E-R 图", 14.5)

    add_heading(document, "6.6.2 一致性、版本与索引", 3)
    add_body(document, "用户角色、订单优先级与状态、车辆类型与状态、坐标来源与状态均使用检查约束限制取值。订单引用网点、任务引用执行方案、坐标修订引用校准人员；删除策略按业务需要使用 RESTRICT、CASCADE 或 SET NULL。高频查询在状态、创建时间、角色、冷链属性、优先级和外键上建立组合索引，历史方案使用时间游标避免深分页 OFFSET。")
    add_body(document, "场景版本激活时使用事务保证同一时刻只有一个 active 版本；优化运行保存场景标识、空间快照、道路元数据和输出指纹。AI 与地图配置使用唯一启用约束保证单一默认连接。数据恢复只重置标准业务域，不覆盖连接配置和账号域，确保使用者已经保存的 Key 和权限不会丢失。")

    add_heading(document, "6.7 部署架构与运维机制", 2)
    add_heading(document, "6.7.1 容器化部署", 3)
    add_body(document, "生产式本地部署由 Docker Compose 编排 frontend、api、worker、db、redis 和 migrate 六类服务。Nginx 暴露统一访问地址，API 与 Worker 使用非超级用户 qizhi_app 连接 PostgreSQL；Redis 开启持久化，用于队列、任务结果和短期缓存；PostgreSQL 与 Redis 数据分别进入 Docker Volume，重新构建镜像不会删除业务数据。")
    add_table(
        document,
        "表 6-7 容器服务与职责",
        ["服务", "技术", "职责", "持久化"],
        [
            ["frontend", "Nginx", "静态站点与 /api 代理", "镜像内静态资源"],
            ["api", "FastAPI", "认证、业务接口、同步求解", "写入 PostgreSQL"],
            ["worker", "Celery", "异步优化任务", "结果回写数据库"],
            ["db", "PostgreSQL/PostGIS", "核心业务与空间数据", "postgres_data"],
            ["redis", "Redis", "队列、结果和缓存", "redis_data"],
            ["migrate", "Alembic", "启动前数据库升级", "执行完成后退出"],
        ],
        [2.4, 3.3, 5.5, 3.45],
    )

    add_heading(document, "6.7.2 一键启动与交付", 3)
    add_body(document, "交付目录提供“启动祁智运平台”“停止祁智运平台”“检查平台状态”和“一键恢复标准数据”四个中文脚本。首次启动会检查 Docker、构建镜像、执行迁移并等待健康检查，完成后使用浏览器访问 http://127.0.0.1。模型服务和地图服务在系统设置中配置并保存在数据库，部署人员无需安装 Python、Node.js、PostgreSQL 或 Redis，也无需手工编辑源代码。")

    add_heading(document, "6.7.3 健康检查、备份与恢复", 3)
    add_body(document, "API 提供存活和就绪检查，分别验证进程、数据库、Redis 与迁移状态；Docker 对数据库、Redis、API 和前端设置依赖与健康条件。日常故障定位先执行状态检查，再查看 API 与 Worker 日志。数据库使用 pg_dump 自定义格式备份，恢复前暂停写入服务并在独立环境验证。")
    add_body(document, "标准数据恢复会在执行前生成可追踪备份，之后恢复固定订单、车辆、网点、算法实验和故障事件，使每次业务操作得到稳定结果。该流程不执行 docker compose down -v，也不清理数据卷；涉及删除数据卷或更换加密密钥的操作必须由管理员明确确认。")
    add_heading(document, "6.7.4 开发阶段规划", 3)
    add_body(document, "平台开发以数据可用、求解可用、业务闭环、智能增强和稳定交付为连续里程碑。测试、安全审计与文档同步贯穿全部阶段，使功能扩展与交付材料保持一致。")

    add_heading(document, "6.8 模式创新与应用价值", 2)
    add_heading(document, "6.8.1 从“路线展示”到“可执行闭环”", 3)
    add_body(document, "传统物流可视化往往在生成路线后结束，祁智运则把方案继续推进到审批、任务拆分、驾驶员确认和完成归档。路线不只是一条地图折线，而是包含车辆、停靠顺序、容量、时间窗、冷链、成本与执行状态的业务对象。动态事件发生后，系统以已执行快照为边界重新求解，从而实现计划与执行的闭环。")

    add_heading(document, "6.8.2 确定性优化与生成式 AI 的责任分离", 3)
    add_body(document, "平台没有把“接入大模型”包装成路线求解本身，而是把 AI 放在擅长的目标理解、策略表达和候选复核位置，把安全与可行性留给 OR-Tools。透明评分公开每一项权重和依据，AI 不可用时结果仍可复现。这一责任分离既发挥自然语言模型的解释能力，也避免幻觉直接改变车辆任务。")

    add_heading(document, "6.8.3 县域空间数据治理", 3)
    add_body(document, "平台同时治理行政边界、业务网点、坐标修订、道路矩阵和空间快照。高德地图负责用户可见的真实底图，OSRM 负责求解矩阵，离线边界保障弱网可用；每次坐标修改和矩阵降级都可追溯。相比仅在前端调整标点，这种数据治理能保证地图、算法和历史结果使用一致的空间事实。")

    add_heading(document, "6.8.4 便捷部署与推广能力", 3)
    add_body(document, "系统把环境安装、数据库迁移、服务启动、配置保存、标准数据恢复和操作引导一起纳入产品。使用人员可以通过中文脚本启动平台，通过网页管理 Key，通过引导完成主流程，通过恢复功能回到标准状态。该能力降低了系统从开发环境转移到其他计算机后的使用风险。")
    add_heading(document, "6.8.5 务实运营组织模式", 3)
    add_body(document, "平台采用县域运营中心统筹、专业小组协同、县乡村服务网络执行、生态伙伴受控接入的组织方式。运营责任与系统角色相互对应，日计划、动态协调、执行反馈和数据复盘构成常态化运营闭环。")
    add_heading(document, "6.8.6 盈利模式重心", 3)
    add_body(document, "商业化路径以可量化的降本增效为前提，由平台订阅、调度优化、实施与数据治理、运营分析构成主要收入支柱，并通过生态接口和增值保障服务延伸。相关收费模式仍需结合实际客户规模、服务边界和采购方式进一步验证。")

    add_heading(document, "6.9 风险边界与后续扩展", 2)
    add_heading(document, "6.9.1 数据边界", 3)
    add_body(document, "标准场景使用固定的祁东县业务样本和经过校准的网点坐标，适合验证流程、算法约束和系统稳定性，但不应直接替代企业真实历史订单。需求预测在样本不足时明确使用统计基线；成本结果来自车辆单位里程成本和道路矩阵，未覆盖过路费、临时装卸费、司机绩效等全部生产成本。")

    add_heading(document, "6.9.2 外部服务风险", 3)
    add_body(document, "OSRM、高德地图和大模型均属于外部依赖，可能出现额度、白名单、网络、鉴权和服务变更。平台通过连接测试、超时、缓存、离线边界和确定性降级保持核心业务可用，但真实生产部署仍需申请正式配额、配置合法域名、监控调用失败率，并制定 Key 轮换方案。")

    add_heading(document, "6.9.3 安全与合规风险", 3)
    add_body(document, "正式使用时需修改默认管理员密码、数据库密码和 JWT 密钥，妥善保存 AI_CONFIG_ENCRYPTION_KEY，限制数据库与容器端口暴露。业务订单可能包含联系人、电话与地址，导出、备份和日志应执行最小必要原则，不把真实数据和密钥提交到代码仓库。公开注册可按实际组织要求扩展为审核后启用或邀请码机制。")

    add_table(
        document,
        "表 6-8 核心风险与应对策略",
        ["风险类别", "主要风险", "应对策略"],
        [
            ["业务数据", "标准样本与真实业务存在差异", "导入企业数据后重新校准坐标、成本和基线"],
            ["道路服务", "OSRM 或地图服务不可用", "缓存、离线边界、估算矩阵和降级提示"],
            ["模型服务", "鉴权失败、超时或输出不合规", "连接测试、白名单解析、超时和确定性回退"],
            ["账号安全", "弱密码、撞库或越权", "强密码、失败锁定、JWT、RBAC 与审计"],
            ["密钥管理", "密钥泄露或错误轮换", "加密入库、掩码输出、受控环境变量和轮换流程"],
            ["运行交付", "部署人员不熟悉环境配置或误操作", "Docker、一键脚本、操作引导、备份与标准恢复"],
        ],
        [3.0, 5.1, 6.55],
    )
    add_body(document, "综上，祁智运已经形成从基础数据、空间路网、约束优化、AI 协同、动态事件到审批执行和运维恢复的完整系统链路。其价值不在于堆叠算法名词，而在于把县域物流中可以量化的约束、可以复核的决策和可以执行的任务统一到一个稳定、可移交的平台中。")

    # Chapter 7 — delivery, verification and acceptance.
    add_heading(document, "7  交付、测试与验收", 1)
    add_heading(document, "7.1 交付清单", 2)
    add_table(document, "表 7-1 平台交付物", ["交付物", "用途", "验收方式"], [
        ["祁智运县域物流智能决策平台代码", "前后端、迁移、数据和脚本", "按部署说明启动并检查健康状态"],
        ["Docker Compose 配置", "数据库、缓存、API、Worker、前端编排", "docker compose ps 全部健康"],
        ["中文启动/停止/检查/恢复脚本", "简化部署与运行操作", "双击脚本后浏览器可访问"],
        ["平台设计与实现文档", "方案、实现、数据依据与验收证据", "逐章核对截图、表格与平台页面"],
        ["图文使用手册 HTML/DOCX", "面向业务操作人员的说明", "按流程完成一次闭环"],
        ["截图与架构图资源", "页面证据和汇报材料", "文件可打开，图注与页面一致"],
    ], [3.3, 6.0, 4.8])
    add_heading(document, "7.2 测试策略与结果", 2)
    add_body(document, "测试分为后端接口契约、算法与数据一致性、前端构建、容器健康和浏览器端到端五层。接口测试覆盖注册登录、权限、场景、调度、事件、执行、配置、审计和参考数据摘要；浏览器截图脚本覆盖登录、注册、总览、调度、事件、执行、实验、数据、服务设置、准备度、场景和历史等页面。每次测试均使用稳定场景，避免随机数据造成误判。")
    add_table(document, "表 7-2 验收场景与判定条件", ["编号", "场景", "通过条件"], [
        ["A-01", "首次启动", "数据库迁移成功，前端首页和存活/就绪检查返回 200"],
        ["A-02", "注册与登录", "弱密码拒绝；注册运营用户；登录成功并按角色进入页面"],
        ["A-03", "生成调度", "16 条订单进入可行方案；路线、里程、成本、车辆数和服务率可见"],
        ["A-04", "动态故障", "冻结已执行任务，故障车辆退出剩余求解，处置状态与消息同步更新"],
        ["A-05", "审批执行", "方案提交、审批、下发、司机接单/发车/完成状态合法推进"],
        ["A-06", "AI/地图配置", "Key 掩码保存；测试结果可见；鉴权失败不泄漏密钥"],
        ["A-07", "恢复基线", "业务数据恢复到固定样本，账号和外部服务配置保留"],
        ["A-08", "审计追溯", "登录、配置、求解、事件和执行操作可在审计/历史页查询"],
    ], [1.5, 4.2, 8.4])
    add_callout(document, "质量复核", "验收应同时检查接口契约、页面流程、约束可行性、外部服务状态和数据恢复结果；任何一项不满足，都应保留运行编号、日志和复现步骤。", fill="#FFF5DD", accent="#B97915")
    add_heading(document, "7.3 标准运行流程", 2)
    add_numbered(document, [
        "双击“启动祁智运平台”，等待状态检查提示服务正常。",
        "浏览器打开本机地址，使用管理员或已创建的运营账号登录。",
        "先看运营总览和场景页，确认订单、车辆、路网来源及服务状态。",
        "进入智能调度生成方案，查看路线卡片、评分和约束说明后提交审批。",
        "如需验证异常处置，进入动态事件处理固定故障，再查看执行中心中的新任务。",
        "操作完成或数据被改乱时，双击“一键恢复标准场景数据”，确认恢复提示后重新开始。",
    ])
    add_figure(document, "13-deployment-topology.png", "图 7-1 交付环境的 Docker Compose 拓扑", 14.0)
    add_figure(document, "25-development-roadmap-image2.png", "图 7-2 平台开发阶段规划图", 14.0)
    add_heading(document, "7.4 运维组织与后续迭代", 2)
    add_body(document, "日常运营由县域运营中心维护场景和计划，调度人员处理方案与异常，驾驶员反馈任务状态，系统管理员维护账号、密钥、地图和审计。后续迭代应优先补充经核验的站点坐标、真实历史订单、过路费与装卸成本、更多车辆班次和生产级监控，再考虑多租户、移动端和更复杂的跨县干线协同。")
    add_figure(document, "26-operations-organization-image2.png", "图 7-3 平台务实运营组织模式图", 14.0)
    add_heading(document, "7.5 应用价值与商业化边界", 2)
    add_body(document, "平台的可量化价值集中在三方面：将研究资料转化为统一的县域空间底座；将经验调度转化为满足硬约束的可行路线；将路线继续推进为可审批、可下发、可复盘的执行任务。平台盈利模式图用于讨论订阅、实施与数据治理、运营分析等可能路径，不代表已有客户合同或经营收入。")
    add_figure(document, "27-profit-model-image2.png", "图 7-4 平台价值与盈利模式重心架构图", 14.0)

    # Appendices — practical references for implementation and acceptance.
    add_heading(document, "附录 A  稳定场景数据字典", 1)
    add_table(document, "表 A-1 订单、车辆与节点核心字段", ["对象", "字段", "说明"], [
        ["订单", "order_id", "稳定订单编号 O-001～O-016"],
        ["订单", "weight_kg", "重量；参与容量约束与需求汇总"],
        ["订单", "node_id", "服务节点；必须存在于标准场景"],
        ["订单", "cold_chain / priority", "冷链标记与优先级"],
        ["订单", "time_window", "服务时间窗；参与时间可行性校验"],
        ["车辆", "vehicle_id", "湘 D·A1234 等稳定车辆编号"],
        ["车辆", "capacity_kg", "额定容量；不可超载"],
        ["车辆", "cold_chain / available", "冷链能力与可用状态"],
        ["车辆", "cost_per_km", "单位里程成本；参与透明评分"],
        ["节点", "node_id / name", "配送中心、乡镇或村级服务节点"],
        ["节点", "longitude / latitude", "校准后的地图坐标"],
        ["节点", "coordinate_source", "坐标来源、精度与校准状态"],
    ], [2.2, 4.0, 8.0])
    add_heading(document, "附录 B  API 与页面对应关系", 1)
    add_table(document, "表 B-1 核心 API 组", ["API 组", "主要能力", "前端入口"], [
        ["auth", "注册、登录、当前用户、登出", "登录/注册页"],
        ["analytics", "预测、乡镇排行、风险和路网状态", "运营总览"],
        ["scenarios / master-data", "场景、订单、车辆、节点与校准", "场景/基础数据"],
        ["dispatch / runs", "生成、评分、提交与历史运行", "智能调度/方案历史"],
        ["events", "事件识别、处置和增量重调度", "动态事件"],
        ["execution", "审批、下发、司机任务与完成", "执行中心"],
        ["experiments", "算法对照、压力场景与指标", "算法评估"],
        ["settings", "AI、地图、用户、安全和审计", "系统设置"],
        ["reference", "祁东设施、需求、候选和数据质量摘要", "运营解释层/接口"],
        ["health", "存活、就绪和依赖检查", "启动脚本/运维"],
    ], [3.4, 6.6, 4.2])
    add_heading(document, "附录 C  配置与安全清单", 1)
    add_table(document, "表 C-1 运行配置的安全处理", ["配置项", "保存位置/方式", "交付注意事项"], [
        ["AI 服务 Key", "系统设置中加密入库，界面首尾掩码", "不要写入代码、截图、日志或 Git"],
        ["高德 Web/Security Key", "地图连接加密入库，安全代理服务端追加", "配置域名白名单和调用额度"],
        ["JWT 密钥", "环境变量注入", "生产环境使用随机长密钥并轮换"],
        ["数据库密码", "环境变量/本机配置文件", "不要使用示例值，不开放数据库端口"],
        ["OSRM 地址", "环境变量或路由设置", "记录真实路网与降级矩阵状态"],
        ["备份文件", "本机受控目录", "限制访问，恢复前先验证备份可读"],
    ], [3.1, 6.0, 5.1])
    add_heading(document, "附录 D  参考资料版本与引用边界", 1)
    add_body(document, "平台使用的资料版本包括：第四章祁东县县乡村三级物流设施建设相关概况（优化版）、第一版：祁东县城乡物流设施体系研究、第一部分(1)、系统摸排祁东县物流行业整体发展水平、政策梳理、第三章基于 AHP-GIS 耦合模型的乡镇仓储分拨站选址适宜性前置筛选及其第一版、县级设施/乡镇站点/人口密度/特色产业/乡镇需求与成本/农村寄递服务站等表格，以及实时路径规划优化系统代码与手册 PDF。")
    add_body(document, "引用原则是“资料支撑场景，平台支撑实现”：研究资料中的规模、设施和政策用于解释为什么设计这些模块；平台代码、API、数据库和截图用于证明已经实现哪些能力。历史版本的 60 分钟服务圈、未经核验的站点逐点坐标和参考项目的 Rails/Mapbox 架构不作为平台运行时事实。")
    add_heading(document, "附录 E  图表与截图索引", 1)
    add_table(document, "表 E-1 主要图表清单", ["编号范围", "内容", "来源"], [
        ["图 0-1", "祁东县县域物流三级配送网络示意", "项目空间示意图；用于场景说明"],
        ["图 3-1～图 3-3", "平台逻辑、功能框架和业务闭环", "Image2 生成图与代码绘图"],
        ["图 4-1～图 4-16", "登录、注册、总览、调度、事件、执行、实验、设置等页面", "平台运行实例 Playwright 截图"],
        ["图 7-1～图 7-4", "部署拓扑、开发路线、运营组织、价值模式", "代码绘图与 Image2 参考图"],
        ["表 2-2～表 2-5", "祁东县设施、产业、AHP-GIS 与候选排序", "新增资料核对与运行时摘要"],
        ["表 7-1～表 7-2", "交付物和验收场景", "平台代码与脚本对应关系"],
    ], [3.0, 7.0, 4.2])
    add_source_note(document, "截图采集说明：页面截图在平台运行实例中生成，采集脚本位于 scripts/capture_project_document.cjs；外部天气、地图和路网服务的实时状态可能随网络与配额变化。")

    configure_header_footer(document)
    document.core_properties.title = "祁智运县域物流智能决策平台设计与实现"
    document.core_properties.subject = "县域物流智能调度、路径规划、可视化与执行闭环"
    document.core_properties.author = "祁智运项目组"
    document.core_properties.keywords = "县域物流, OR-Tools, OSRM, AI协同, 动态重调度, Docker"
    renumber_figures(document)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
