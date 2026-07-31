from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


GREEN = "0D5138"
GREEN_2 = "15734A"
MINT = "EAF4EF"
YELLOW = "F5B83B"
INK = "17372B"
MUTED = "5E7066"
LINE = "D8E3DD"
PALE = "F5F8F6"
WHITE = "FFFFFF"
RED = "C94B43"


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, color: str = LINE, size: int = 6) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=140, bottom=110, end=140) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeatable_heading(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep_next = OxmlElement("w:keepNext")
    keep_lines = OxmlElement("w:keepLines")
    p_pr.append(keep_next)
    p_pr.append(keep_lines)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    begin_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(begin)
    instr_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    instr_run._r.append(instr)
    separate_run = paragraph.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run._r.append(separate)
    result_run = paragraph.add_run("1")
    result_run.font.size = Pt(9)
    result_run.font.color.rgb = RGBColor.from_string(MUTED)
    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)
    run = paragraph.add_run(" 页")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Title", 30, WHITE, 0, 12),
        ("Subtitle", 13, WHITE, 0, 10),
        ("Heading 1", 16, GREEN, 18, 10),
        ("Heading 2", 13, GREEN_2, 14, 7),
        ("Heading 3", 11.5, INK, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)


def add_header_footer(section) -> None:
    section.different_first_page_header_footer = True
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.text = "祁智运 · 县域物流智能决策平台  /  图文使用手册"
    paragraph.style = "Caption"
    paragraph.runs[0].font.color.rgb = RGBColor.from_string(GREEN)
    paragraph.runs[0].font.bold = True
    paragraph.runs[0].font.size = Pt(8.5)
    footer = section.footer
    add_page_field(footer.paragraphs[0])


def add_section_title(document: Document, number: str, title: str, subtitle: str) -> None:
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(1.3)
    table.columns[1].width = Cm(14.8)
    left, right = table.rows[0].cells
    left.width = Cm(1.3)
    right.width = Cm(14.8)
    set_cell_shading(left, YELLOW)
    set_cell_shading(right, MINT)
    for cell in (left, right):
        set_cell_border(cell, WHITE, 0)
        set_cell_margins(cell, 120, 140, 120, 140)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(number)
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(GREEN)
    p = right.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor.from_string(GREEN)
    p = right.add_paragraph(subtitle)
    p.paragraph_format.space_after = Pt(0)
    p.runs[0].font.size = Pt(9)
    p.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_callout(document: Document, title: str, body: str, kind: str = "info") -> None:
    colors = {
        "info": (MINT, GREEN_2),
        "warning": ("FFF5DF", "A86A08"),
        "danger": ("FDECEB", RED),
        "neutral": (PALE, MUTED),
    }
    fill, accent = colors[kind]
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, accent, 8)
    set_cell_margins(cell, 130, 180, 130, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(title)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(accent)
    run.font.size = Pt(10.5)
    p = cell.add_paragraph(body)
    p.paragraph_format.space_after = Pt(0)
    p.runs[0].font.size = Pt(9.5)
    p.runs[0].font.color.rgb = RGBColor.from_string(INK)


def new_list_number_id(document: Document) -> int:
    numbering = document.part.numbering_part.element
    style_num_id = int(
        document.styles["List Number"]._element.pPr.numPr.numId.val
    )
    abstract_id = "0"
    for num in numbering.findall(qn("w:num")):
        if int(num.get(qn("w:numId"))) == style_num_id:
            abstract = num.find(qn("w:abstractNumId"))
            if abstract is not None:
                abstract_id = abstract.get(qn("w:val"))
            break
    existing = [
        int(num.get(qn("w:numId")))
        for num in numbering.findall(qn("w:num"))
    ]
    num_id = max(existing, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return num_id


def add_steps(document: Document, steps: list[str]) -> None:
    num_id = new_list_number_id(document)
    for text in steps:
        p = document.add_paragraph(style="List Number")
        p_pr = p._p.get_or_add_pPr()
        num_pr = p_pr.find(qn("w:numPr"))
        if num_pr is None:
            num_pr = OxmlElement("w:numPr")
            p_pr.append(num_pr)
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id_node = OxmlElement("w:numId")
        num_id_node.set(qn("w:val"), str(num_id))
        num_pr.append(ilvl)
        num_pr.append(num_id_node)
        p.add_run(text)


def add_bullets(document: Document, items: list[str]) -> None:
    for text in items:
        p = document.add_paragraph(style="List Bullet")
        p.add_run(text)


def add_screenshot(document: Document, path: Path, caption: str, width=6.55) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    picture = run.add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption.split("  ", 1)[0])
    caption_p = document.add_paragraph()
    caption_p.style = "Caption"
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.keep_with_next = True
    caption_p.paragraph_format.space_after = Pt(7)
    r = caption_p.add_run(caption)
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(MUTED)


def add_page_break(document: Document) -> None:
    document.add_page_break()


def add_key_value_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(4.0)
    table.columns[1].width = Cm(12.0)
    header = table.rows[0]
    header.cells[0].text = "项目"
    header.cells[1].text = "标准值 / 说明"
    set_repeat_table_header(header)
    for cell in header.cells:
        set_cell_shading(cell, GREEN)
        set_cell_border(cell, WHITE, 4)
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(WHITE)
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
        for index, cell in enumerate(cells):
            set_cell_shading(cell, PALE if index == 0 else WHITE)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if index == 0:
                cell.paragraphs[0].runs[0].font.bold = True
                cell.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(GREEN)


def build_manual(screens: Path, output: Path) -> None:
    document = Document()
    configure_styles(document)
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    add_header_footer(section)

    cover = document.add_table(rows=1, cols=1)
    cover_cell = cover.cell(0, 0)
    set_cell_shading(cover_cell, GREEN)
    set_cell_border(cover_cell, GREEN, 0)
    set_cell_margins(cover_cell, 550, 480, 550, 480)
    p = cover_cell.paragraphs[0]
    r = p.add_run("祁 智 运")
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(YELLOW)
    p = cover_cell.add_paragraph("县域物流智能决策平台")
    p.style = "Title"
    p = cover_cell.add_paragraph("图文使用手册")
    p.style = "Title"
    p.runs[0].font.color.rgb = RGBColor.from_string(YELLOW)
    p = cover_cell.add_paragraph("从启动、调度到故障处置与驾驶员执行")
    p.style = "Subtitle"
    p = cover_cell.add_paragraph()
    p.paragraph_format.space_before = Pt(70)
    r = p.add_run("标准运营基线 v1")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(WHITE)
    p = cover_cell.add_paragraph("12 个固定订单  ·  5 辆固定车辆  ·  固定路网快照  ·  固定故障事件")
    p.runs[0].font.size = Pt(10)
    p.runs[0].font.color.rgb = RGBColor.from_string("D8EEE3")
    p = document.add_paragraph("版本：2026.07  |  适用地址：http://127.0.0.1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    p.runs[0].font.size = Pt(9)

    add_page_break(document)
    add_section_title(document, "00", "先看这一页", "不懂技术也可以按顺序完成一次完整操作")
    add_callout(
        document,
        "最短操作路径",
        "双击“启动祁智运平台.bat” → 浏览器打开平台 → 管理员登录 → 运营总览 → 生成今日调度 → 动态事件 → 启动处置 → 执行中心。",
    )
    document.add_heading("账号与入口", level=2)
    add_key_value_table(
        document,
        [
            ("平台地址", "http://127.0.0.1"),
            ("管理员账号", "admin"),
            ("管理员密码", "QizhiDemo2026!"),
            ("驾驶员账号", "driver"),
            ("驾驶员密码", "DriverDemo2026!"),
        ],
    )
    document.add_heading("这套平台解决什么问题", level=2)
    add_bullets(
        document,
        [
            "把县域订单、车辆容量、冷链要求、时间窗和班次约束放入同一个调度模型。",
            "使用 OR-Tools 生成可执行路线，地图展示真实路网走向、里程和任务顺序。",
            "车辆故障时冻结已完成部分，只对剩余任务重新分配并记录处置依据。",
            "管理员负责全局调度与审批，驾驶员只查看和执行分配给自己的任务。",
        ],
    )
    add_callout(
        document,
        "稳定结果说明",
        "标准运营基线已固定订单、车辆、OSRM 路网矩阵和故障事件。AI 可以解释结果，但不会改变这套标准路线，因此重复操作不会出现方案漂移。",
        "warning",
    )

    add_page_break(document)
    add_section_title(document, "01", "启动、检查与关闭", "四个批处理文件已把数据库和 Docker 操作封装好")
    add_key_value_table(
        document,
        [
            ("启动祁智运平台.bat", "自动启动 Docker Desktop、数据库、后端和页面，并打开浏览器。"),
            ("检查平台状态.bat", "检查数据库、缓存、调度服务和页面是否正常。"),
            ("停止祁智运平台.bat", "停止平台服务；业务数据与密钥配置不会丢失。"),
            ("一键恢复标准演示数据.bat", "备份当前数据库后恢复标准订单、车辆和事件；保留 AI 与高德配置。"),
        ],
    )
    document.add_heading("首次使用", level=2)
    add_steps(
        document,
        [
            "确认电脑已安装 Docker Desktop；Python、Node.js 和数据库不需要另外安装。",
            "双击“启动祁智运平台.bat”，等待窗口出现“全部服务已经就绪”。",
            "浏览器若未自动打开，手动访问 http://127.0.0.1。",
            "出现异常时先双击“检查平台状态.bat”，再按手册末尾的排查表处理。",
        ],
    )
    add_callout(
        document,
        "不要直接关闭 Docker 数据目录",
        "正常结束请使用“停止祁智运平台.bat”。直接删除 Docker 容器或数据卷可能造成业务数据不可恢复。",
        "danger",
    )

    add_page_break(document)
    add_section_title(document, "02", "登录平台", "管理员与驾驶员共用登录页，登录后按角色展示权限")
    add_screenshot(document, screens / "01-login.png", "图 1  企业化登录页：输入账号和密码后点击“安全登录”")
    add_steps(
        document,
        [
            "管理员使用 admin 登录，负责调度、事件处置、审批、配置与数据治理。",
            "驾驶员使用 driver 登录，仅查看与执行本人被分配的任务。",
            "连续 5 次输入错误会锁定 15 分钟；不要反复尝试不确定的密码。",
        ],
    )

    add_page_break(document)
    add_section_title(document, "03", "跟随内置操作引导", "第一次登录会高亮下一步入口，随时可在顶栏重新打开")
    add_screenshot(document, screens / "02-operation-guide.png", "图 2  操作引导会遮罩无关区域，并指出当前应点击的位置")
    add_bullets(
        document,
        [
            "引导共 5 步：运营总览、生成调度、智能调度、动态事件、执行中心。",
            "引导只切换页面，不会自动重算、审批或下发任务。",
            "想中途退出可点击右上角 ×；再次使用点击顶栏“操作引导”。",
        ],
    )

    add_page_break(document)
    add_section_title(document, "04", "读懂运营总览", "先确认需求、运力、路网与风险，再进入调度")
    add_screenshot(document, screens / "03-overview.png", "图 3  运营总览：需求预测、路网态势、运力建议和风险预警集中在一页")
    add_key_value_table(
        document,
        [
            ("今日预测需求", "固定业务样本的当日需求规模，用于判断运力是否充足。"),
            ("建议运力", "按目标装载率与冷链需求计算的建议车辆数。"),
            ("路网矩阵", "标准基线显示 OSRM；表示里程与时间来自固定真实路网快照。"),
            ("预测准时率", "当前容量和时间窗约束下的可达成比例。"),
        ],
    )

    add_page_break(document)
    add_section_title(document, "05", "生成固定调度方案", "标准基线始终使用同一批订单、车辆、路网与求解参数")
    add_screenshot(document, screens / "04-fixed-dispatch.png", "图 4  智能调度中心：地图、路线任务和优化前后指标同步展示", width=6.15)
    add_steps(
        document,
        [
            "从运营总览点击“生成今日调度”，或在左侧进入“智能调度”。",
            "点击右上角“重新生成方案”，等待按钮恢复可用。",
            "点击右侧车辆卡片切换路线；地图上的编号即该车的访问顺序。",
            "查看综合绩效评分、订单量、货量、准时率和优化前后对比。",
        ],
    )
    add_callout(
        document,
        "标准结果识别",
        "标准基线应为 12 单、5 辆可用车，最终使用 3 辆车，订单服务率 100%。总里程约 339 km、预计成本约 1,196 元；页面缓存命中与求解耗时可变化，但路线和核心指标不变。",
    )

    add_page_break(document)
    add_section_title(document, "06", "处理固定车辆故障", "事件车辆、位置和紧急订单均已固定，便于重复操作与复核")
    add_screenshot(document, screens / "06-event-after.png", "图 5  故障处置完成：湘D·A1234 在尚义村的事件已闭环，并生成替代方案")
    add_steps(
        document,
        [
            "进入“动态事件”，点击右上角“启动处置”。",
            "系统固定识别 V-COLD-01（湘D·A1234）在尚义村发生故障。",
            "系统冻结已完成节点，加入固定紧急订单，并启用固定应急冷藏备用车。",
            "确认右侧四个环节均为完成：事件识别、约束重建、增量求解、方案下发。",
            "查看“处置结果”和“优化决策依据”，确认冷链约束为 100% 满足。",
        ],
    )
    add_callout(
        document,
        "图标何时消失",
        "处置完成后，故障点不再作为待处理告警保留；顶部状态改为“事件闭环”。若仍显示待处理红点，刷新页面或执行一次标准数据恢复。",
        "warning",
    )

    add_page_break(document)
    add_section_title(document, "07", "审批与执行闭环", "管理员确认方案后，驾驶员接收并更新任务状态")
    add_screenshot(document, screens / "07-execution.png", "图 6  执行中心：管理员查看方案状态、车辆任务与操作记录")
    add_bullets(
        document,
        [
            "管理员：查看调度方案 → 审核约束与路线 → 批准并下发 → 监控执行。",
            "驾驶员：登录后只看到本人任务 → 接单 → 开始执行 → 逐站完成 → 完成整条路线。",
            "每次审批、下发和状态更新都会形成记录，便于追溯。",
        ],
    )
    add_callout(
        document,
        "角色界面为什么不同",
        "这是权限隔离，不是数据错误。管理员看到县域全局路网与全部车辆；驾驶员只看到本人被授权的路线和任务，避免越权访问其他车辆数据。",
        "info",
    )

    add_page_break(document)
    add_section_title(document, "08", "驾驶员视角", "驾驶员不参与全局求解，只执行管理员已下发的任务")
    add_screenshot(document, screens / "11-driver-execution.png", "图 7  驾驶员执行中心：未下发时显示空任务；下发后仅展示本人路线", width=6.0)
    add_steps(
        document,
        [
            "使用 driver / DriverDemo2026! 登录。",
            "进入“执行中心”，确认车牌、节点顺序、预计到达时间和冷链标识。",
            "按实际进度执行接单、出发、到站和完成操作。",
            "如出现车辆故障，停止继续操作并通知管理员在“动态事件”中处置。",
        ],
    )

    add_page_break(document)
    add_section_title(document, "09", "AI 服务配置", "密钥只在新增或轮换时录入，保存后仅显示掩码")
    add_screenshot(document, screens / "08-ai-settings.png", "图 8  AI 服务连接列表：默认启用连接、测试状态和更新人清晰分离")
    add_steps(
        document,
        [
            "管理员进入“系统设置 → AI 服务”，点击“新建连接”。",
            "填写配置名称、服务商、Base URL、模型 ID 和 API Key。",
            "保存后点击“测试”；只有测试正常的连接才设置为默认启用。",
            "编辑已有连接时 API Key 留空表示不修改，页面不会返回明文。",
        ],
    )
    add_callout(
        document,
        "AI 在标准基线中的边界",
        "OR-Tools 负责满足容量、冷链、时间窗和班次硬约束并生成路线。AI 用于解释、风险提示和非标准业务的策略建议；标准运营基线中 AI 不改变最终路线。",
        "warning",
    )

    add_page_break(document)
    add_section_title(document, "10", "高德地图配置", "Web Key 与安全密钥已加密保存，一键恢复不会清除")
    add_screenshot(document, screens / "09-map-settings.png", "图 9  地图服务配置：连接以列表方式管理，密钥仅显示掩码")
    add_steps(
        document,
        [
            "管理员进入“系统设置 → 地图服务”。",
            "新增或编辑高德连接，填写 Web 端（JS API）Key 与配套安全密钥。",
            "保存并测试连接，然后设为默认启用。",
            "地图显示路线使用 GCJ-02；调度矩阵使用固定 OSRM 快照，两者职责不同。",
        ],
    )
    add_callout(
        document,
        "配置会保留",
        "正常停止、重新启动以及“一键恢复标准演示数据”都会保留 AI Key 和高德配置。只有主动删除连接或删除 Docker 数据卷才会丢失。",
        "info",
    )

    add_page_break(document)
    add_section_title(document, "11", "一键恢复标准状态", "操作错、数据改乱或事件残留时，恢复到可重复的标准起点")
    add_steps(
        document,
        [
            "关闭正在操作的平台页面，双击“一键恢复标准演示数据.bat”。",
            "确认恢复；脚本会先自动备份当前 PostgreSQL 数据库。",
            "等待窗口显示“标准演示数据已经恢复”。",
            "重新打开 http://127.0.0.1，并使用标准账号登录。",
        ],
    )
    add_key_value_table(
        document,
        [
            ("恢复内容", "18 个已校准节点、12 个订单、5 辆车、标准账号、初始调度结果、待处置事件状态。"),
            ("清理内容", "已产生的审批、执行、历史运行和临时事件记录。"),
            ("保留内容", "AI 服务连接、高德地图连接、数据质量规则。"),
            ("备份位置", "项目目录下 backups 文件夹，文件名包含恢复时间。"),
        ],
    )
    add_callout(
        document,
        "固定故障事件",
        "恢复后事件重新回到待处置状态：湘D·A1234 / V-COLD-01，尚义村，发生时刻第 90 分钟，固定紧急订单 O-URGENT-01。",
        "neutral",
    )

    add_page_break(document)
    add_section_title(document, "12", "8 分钟讲解顺序", "按业务问题 → 求解 → 事件 → 执行的逻辑说明，避免陷入技术细节")
    add_key_value_table(
        document,
        [
            ("0:00—1:00", "登录与运营总览：说明县域配送存在多订单、多车型、冷链与时间窗约束。"),
            ("1:00—3:30", "生成今日调度：查看 3 条路线、地图顺序和优化前后指标。"),
            ("3:30—5:30", "动态事件：启动固定车辆故障处置，说明冻结已执行任务和增量重算。"),
            ("5:30—6:40", "执行中心：说明管理员审批下发、驾驶员接单和状态回传。"),
            ("6:40—7:30", "AI 与算法边界：OR-Tools 保证可行性，AI 负责解释和风险提示。"),
            ("7:30—8:00", "总结：路线可计算、过程可执行、结果可追溯，且可一键恢复标准状态。"),
        ],
    )
    document.add_heading("一句话概括", level=2)
    add_callout(
        document,
        "平台价值",
        "祁智运把县域物流的订单、运力和突发事件统一成可计算的调度问题，再把求解结果转成地图路线、审批任务和驾驶员执行闭环。",
    )

    add_page_break(document)
    add_section_title(document, "13", "常见问题与快速排查", "先确认服务与账号，再检查配置和业务状态")
    add_key_value_table(
        document,
        [
            ("页面打不开", "双击“检查平台状态.bat”；若 Docker Desktop 未运行，重新双击启动文件。"),
            ("登录失败", "确认大小写和感叹号；连续失败后等待 15 分钟，或由管理员解除锁定。"),
            ("地图空白", "检查“系统设置 → 地图服务”的默认连接与测试状态，然后刷新页面。"),
            ("AI 401", "API Key 无效或 Base URL 与服务商不匹配；编辑连接后重新测试。"),
            ("结果与手册不同", "先执行一键恢复，再确认使用“标准运营基线 v1”生成方案。"),
            ("故障红点仍在", "确认处置进度为 100%；刷新页面，仍异常则执行一键恢复。"),
            ("驾驶员看不到全局地图", "正常权限设计；驾驶员仅能访问本人任务，管理员账号可查看全局。"),
            ("想重新开始", "使用“一键恢复标准演示数据.bat”，AI 与高德配置会继续保留。"),
        ],
    )
    add_callout(
        document,
        "交接建议",
        "把项目整个文件夹、Docker Desktop 安装包、本手册和标准账号一并交付。首次交接时完整走一遍 8 分钟讲解顺序，并让接收者亲自完成一次恢复。",
        "warning",
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--screens", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build_manual(args.screens.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
